import base64
import copy
import io
import threading
import unittest
from types import SimpleNamespace
from unittest.mock import patch
from openpyxl import load_workbook
import overview_export_service as exporter
import test_publication_times as publication


class FakeStore:
    def __init__(self, count=401):
        self.pull_lock = threading.RLock()
        self.lock = threading.RLock()
        self.requests = []
        self.rows = [{"note_id": str(i).zfill(24), "title": f"样本{i}", "is_deleted": i % 2,
                      "analysis_is_negative": "是" if i % 3 else "否"} for i in range(count)]
        self.bad = None

    def _validate_data_overview_read_snapshot(self, token, action):
        if token != "token": raise ValueError("快照已变化")

    def _connect(self):
        return SimpleNamespace(close=lambda: None)

    def query_data_overview(self, payload):
        self.requests.append(copy.deepcopy(payload))
        start = (payload["page"] - 1) * 200
        rows = [{k:r.get(k) for k in payload["fields"]} for r in self.rows[start:start+200]]
        result = dict(ok=True, consistentSnapshot=True, dataset="notes", snapshotToken="token",
                      total=len(self.rows), page=payload["page"], pageSize=200, rows=rows)
        if self.bad: self.bad(result)
        return result


def specs(_db, _dataset):
    return [SimpleNamespace(key=k,label=k,data_type="text",action="")
            for k in ["note_id","title","is_deleted","analysis_is_negative"]]


class FullExportTests(unittest.TestCase):
    def payload(self, count=401):
        return dict(dataset="notes", snapshotToken="token", expectedTotal=count, fields=["title"], search="筛选",
                    filter={"children":[{"field":"title","operator":"contains","value":"样本"}]},
                    sort=[{"field":"note_id","direction":"asc"}], groupThreads=False)

    @patch.object(exporter, "build_field_specs", side_effect=specs)
    def test_401_rows_all_fields_same_filters_three_pages(self, _specs):
        store=FakeStore();payload=self.payload();original=copy.deepcopy(payload)
        result=exporter.export_filtered_workbook(store,payload)
        self.assertEqual(payload,original)
        self.assertEqual([1,2,3],[p["page"] for p in store.requests])
        for req in store.requests:
            self.assertEqual(req["filter"],payload["filter"])
            self.assertEqual(req["search"],payload["search"])
            self.assertIn("analysis_is_negative",req["fields"])
        wb=load_workbook(io.BytesIO(base64.b64decode(result["contentBase64"])))
        self.assertEqual(402,wb["筛选帖子"].max_row)
        self.assertEqual(401,result["total"])
        self.assertEqual("000000000000000000000000",wb["筛选帖子"].cell(2,1).value)

    @patch.object(exporter, "build_field_specs", side_effect=specs)
    def test_changed_total_token_missing_id_duplicate_and_short_page_fail(self, _specs):
        changes=[lambda r:r.update(total=402),lambda r:r.update(snapshotToken="other"),
                 lambda r:r["rows"][0].update(note_id=""),
                 lambda r:r["rows"][1].update(note_id=r["rows"][0]["note_id"]),
                 lambda r:r["rows"].pop()]
        for change in changes:
            with self.subTest(change=change), patch.object(exporter,"build_overview_workbook") as build:
                store=FakeStore();store.bad=change
                with self.assertRaises(ValueError):exporter.export_filtered_workbook(store,self.payload())
                build.assert_not_called()

    def test_more_than_180_fields_split_without_losing_or_misjoining_rows(self):
        large = specs(None, None) + [SimpleNamespace(key=f"extra_{i}",label=f"额外{i}",data_type="text",action="") for i in range(200)]
        store = FakeStore(2)
        with patch.object(exporter,"build_field_specs",return_value=large), patch.object(exporter,"build_overview_workbook",return_value=b"PK\x03\x04") as builder:
            exporter.export_filtered_workbook(store,self.payload(2))
            self.assertEqual(2,len(store.requests))
            self.assertTrue(all(len(p["fields"])<=180 for p in store.requests))
            args=builder.call_args.args
            self.assertEqual(204,len(args[2]))
            self.assertTrue(all(len(row)==204 for row in args[1]))
            self.assertEqual({"main_category":"negative_type","subcategory":"negative_subtype"},args[3]["field_mapping"])
        store = FakeStore(2)
        def reverse_second(result):
            if len(store.requests)==2: result["rows"].reverse()
        store.bad=reverse_second
        with patch.object(exporter,"build_field_specs",return_value=large):
            with self.assertRaisesRegex(ValueError,"关联不一致"):
                exporter.export_filtered_workbook(store,self.payload(2))

    @patch.object(exporter, "build_field_specs", side_effect=specs)
    def test_empty_results_and_stale_snapshot(self, _specs):
        result=exporter.export_filtered_workbook(FakeStore(0),self.payload(0));self.assertEqual(0,result["total"])
        with self.assertRaises(ValueError):exporter.export_filtered_workbook(FakeStore(),{**self.payload(),"snapshotToken":"old"})


class ParentContextTests(unittest.TestCase):
    setUp=publication.TimeStoreTests.setUp
    pull=publication.TimeStoreTests.pull

    def test_filtered_reply_gets_unfiltered_parent_and_all_fields(self):
        self.pull(comments=[{"commentId":"root","content":"非差评一级","commentLevel":1,"author":"父作者"},
                            {"commentId":"reply","parentCommentId":"root","content":"筛选命中","commentLevel":2}])
        schema=self.store.data_overview_schema();self.assertTrue(schema["queryReady"])
        payload=dict(dataset="comments",snapshotToken=schema["snapshotToken"],expectedTotal=1,
                     filter={"children":[{"field":"comment_id","operator":"eq","value":"reply"}]})
        result=self.store.export_data_overview(payload)
        wb=load_workbook(io.BytesIO(base64.b64decode(result["contentBase64"])))
        ws=wb["筛选评论"]
        self.assertEqual(2,ws.max_row)
        self.assertEqual("非差评一级",ws.cell(2,4).value)
        self.assertEqual("二级回复",ws.cell(2,5).value)
        self.assertEqual("reply",ws.cell(2,20).value)
        raw=wb["原始筛选数据"]
        self.assertIn("原始结构化数据",[c.value for c in raw[1]])

    def test_exact_parent_chain_cycle_missing_and_cross_note_are_flagged(self):
        self.pull(comments=[{"commentId":"root","content":"root","commentLevel":1},
                            {"commentId":"mid","parentCommentId":"root","content":"mid","commentLevel":2},
                            {"commentId":"reply","parentCommentId":"mid","content":"reply","commentLevel":2}])
        with self.store._session() as db:
            rows=[dict(db.execute("SELECT * FROM comments WHERE comment_id='reply'").fetchone())]
            exporter.enrich_roots(db,rows);self.assertEqual("root",rows[0]["thread_root_id"])
            db.execute("UPDATE comments SET parent_comment_id='reply' WHERE comment_id='mid'")
            exporter.enrich_roots(db,rows);self.assertEqual("cycle",rows[0]["thread_root_status"])
            db.execute("UPDATE comments SET parent_comment_id='missing-exact' WHERE comment_id='mid'")
            exporter.enrich_roots(db,rows);self.assertEqual("missing",rows[0]["thread_root_status"])
            self.assertIn("missing-exact",rows[0]["thread_root_content"])
            db.execute("UPDATE comments SET note_id='other' WHERE comment_id='mid'")
            exporter.enrich_roots(db,rows);self.assertEqual("cross_note",rows[0]["thread_root_status"])

if __name__=="__main__":unittest.main()
