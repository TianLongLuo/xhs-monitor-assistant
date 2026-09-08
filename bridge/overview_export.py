"""Pure, read-only Excel builder for an already-filtered, complete API snapshot.

No query, paging, deduplication, semantic inference or persistence happens here.
Call build_overview_workbook('comments' | 'notes', rows, columns, context=None).
Rows must contain all snapshot fields, not only the visible columns. The raw sheet
keeps exactly the supplied columns, in input row/column order (including repeats).
The caller owns snapshot/count validation. Context is recorded, never reapplied.

Optional report fields are NOT asserted to exist in the current API schema:
  control: control_level / control / report_control, exactly high or medium;
  translate_comment, main_category, subcategory, highly_similar, judgment_reason.
For another explicitly supplied field, context['field_mapping'] may map one of
those five canonical names to its exact row key. Control mapping is intentionally
not configurable: legacy risk_level/is_negative are never report control levels.
negative_type/negative_subtype are not translated into report categories. Row
highlighting uses analysis_is_negative only, not the legacy AI is_negative flag.

Every input record occupies exactly one main and one raw row, even duplicate or
malformed IDs. Anomalous relationships are isolated, never repaired by text.
ID strings are not trimmed, case-folded, prefixed or canonicalized. Text is stored
as XLSX string cells, never formulas. Unrepresentable/oversized values raise
ValueError with a location, rather than accepting openpyxl's silent truncation.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from collections.abc import Mapping
from datetime import date, datetime
from io import BytesIO
import json
import math
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


COMMENT_HEADERS = (
    '组号', '同组条数', '一级评论者', '一级评论原文', '本条角色', '评论者',
    '评论原文', 'Translate Comment', '控制级别', 'Main Category', 'Subcategory',
    '评论层级', '高度相似', '已删除', '判断依据', '笔记ID', '原帖标题',
    '原帖链接', '帖子作者', '评论ID', '评论时间', '点赞量',
)
COMMENT_WIDTHS = (8, 10, 16, 46, 12, 16, 48, 52, 20, 24, 36, 14, 12, 12,
                  44, 28, 38, 46, 18, 30, 23, 12)
REPORT_FIELDS = {
    'translate_comment': ('translate_comment', 'Translate Comment'),
    'main_category': ('main_category', 'Main Category'),
    'subcategory': ('subcategory', 'Subcategory'),
    'highly_similar': ('highly_similar', '高度相似'),
    'judgment_reason': ('judgment_reason', '判断依据'),
}
CONTROL_FIELDS = ('control_level', 'control', 'report_control')
HEADER = '1C2430'
GROUP_COLORS = ('E7EEF4', 'DCE6EF')
DELETED_COLOR = 'F8D7DA'
NEGATIVE_COLOR = 'E2F0D9'
CATEGORY_COLORS = {'Sales pitch': 'F8D7DA', 'Price discrepancy': 'FFF3CD',
                   'Allergy': 'F3D6E8'}
THIN = Side(style='thin', color='DDDDDD')
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
INVALID_XML = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')
EXCEL_MAX_TEXT = 32767
EXCEL_MAX_ROWS = 1048576
EXCEL_MAX_COLUMNS = 16384


def _boolean(value):
    """Tri-state, deliberately not bool(value); unknown remains None."""
    if value is True or value is False:
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return value == 1
    if isinstance(value, str):
        value = value.strip().lower()
        if value in ('1', 'true', '是'):
            return True
        if value in ('0', 'false', '否'):
            return False
    return None


def _deleted(row, dataset):
    values = [_boolean(row.get('is_deleted'))]
    fields = ['comment_status' if dataset == 'comments' else 'post_status']
    if dataset == 'comments':
        values.append(_boolean(row.get('post__is_deleted')))
        fields.append('post__post_status')
    for field in fields:
        value = row.get(field)
        values.append(True if value in ('已删除', '删除', '已下架') else
                      False if value == '存在' else None)
    # Any explicit deletion wins; absence/unknown is not treated as deleted.
    return True if True in values else False if False in values else None


def _control(row):
    values = [row[k] for k in CONTROL_FIELDS if row.get(k) not in (None, '')]
    if not values:
        return '', '缺少明确专项控制字段'
    if any(value != values[0] for value in values[1:]):
        return '', '专项控制字段冲突'
    if values[0] not in ('high', 'medium'):
        return '', '专项控制字段值未知'
    return values[0], ''


def _report(row, name, mapping):
    keys = (mapping[name],) if name in mapping else REPORT_FIELDS[name]
    for key in keys:
        if row.get(key) is not None and row.get(key) != '':
            return row[key]
    return None


def _text(value):
    if isinstance(value, str):
        return value
    if isinstance(value, (Mapping, list, tuple)):
        return json.dumps(value, ensure_ascii=False, separators=(',', ':'), allow_nan=False)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _check_text(value, location):
    # Count UTF-16 units conservatively, including surrogate pairs for emoji.
    if len(value) > EXCEL_MAX_TEXT or len(value.encode('utf-16-le', errors='surrogatepass')) // 2 > EXCEL_MAX_TEXT:
        raise ValueError(f'{location}: text exceeds Excel 32767-character limit; no data was truncated')
    if INVALID_XML.search(value):
        raise ValueError(f'{location}: text contains an XML-incompatible character')


def _id_key(value):
    # Keep numeric and string source IDs distinct; no alias normalization.
    if isinstance(value, str) and value != '':
        return ('str', value)
    if type(value) is int:
        return ('int', value)
    return None


def _is_id(key):
    return key == 'id' or key.endswith('_id') or key.endswith('ID') or key.endswith('Id')


def _cell(ws, row, col, value, *, key='', data_type=None):
    cell = ws.cell(row, col)
    if value is not None:
        # Strings stay strings even for boolean/numeric column metadata. This
        # preserves original '0', leading zeros, and formula-looking payloads.
        textual = (isinstance(value, (str, Mapping, list, tuple, date, datetime))
                   or _is_id(key) or data_type == 'text'
                   or (type(value) is int and abs(value) >= 10**15))
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError(f'{ws.title}!{cell.coordinate}: non-finite number')
        if textual:
            value = _text(value)
            _check_text(value, f'{ws.title}!{cell.coordinate}')
            cell.value = value
            cell.data_type = 's'  # Explicitly override openpyxl's formula detection.
            cell.number_format = '@'
        elif isinstance(value, (int, float, bool)):
            cell.value = value
        else:
            raise ValueError(f'{ws.title}!{cell.coordinate}: unsupported value type {type(value).__name__}')
    cell.font = Font(name='Calibri', size=11, color='243040')
    cell.alignment = Alignment(vertical='top', wrap_text=True)
    cell.border = GRID
    return cell


def _sheet(wb, title, headers, widths=None):
    if len(headers) > EXCEL_MAX_COLUMNS:
        raise ValueError(f'{title}: exceeds Excel column limit')
    ws = wb.create_sheet(title)
    ws.freeze_panes = 'A2'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.row_dimensions[1].height = 30
    for col, label in enumerate(headers, 1):
        cell = _cell(ws, 1, col, label, data_type='text')
        cell.fill = PatternFill('solid', fgColor=HEADER)
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1] if widths else 24
    return ws


def _finish(ws):
    if ws.max_row > EXCEL_MAX_ROWS:
        raise ValueError(f'{ws.title}: exceeds Excel row limit')
    ws.auto_filter.ref = f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'


def _highlight(ws, excel_row, row, dataset, *, end, start=1, status_col=None, preserve=()):
    deleted = _deleted(row, dataset) is True
    negative = _boolean(row.get('analysis_is_negative')) is True
    color = DELETED_COLOR if deleted else NEGATIVE_COLOR if negative else None
    if color:
        # ws.max_column scans all instantiated cells: do not call it per record
        # in an all-records export (that would make highlighting quadratic).
        for col in range(start, end + 1):
            if col in preserve:
                continue
            cell = ws.cell(excel_row, col)
            cell.fill = PatternFill('solid', fgColor=color)
            cell.font = Font(name='Calibri', size=11, color='243040')
    if deleted and negative:
        # Pink has priority; green remains visible without changing group cells.
        cell = ws.cell(excel_row, status_col or start)
        cell.fill = PatternFill('solid', fgColor=NEGATIVE_COLOR)
        cell.border = Border(left=Side(style='medium', color='548235'),
                             right=THIN, top=THIN, bottom=THIN)


def _flat(wb, title, rows, columns, dataset):
    # Zero-column input still has one physical row per record, explicitly numbered.
    effective = columns or [{'key': '__record_number__', 'label': '筛选记录序号', 'dataType': 'number'}]
    ws = _sheet(wb, title, [c.get('label', c['key']) for c in effective])
    for r, source in enumerate(rows, 2):
        for c, spec in enumerate(effective, 1):
            value = source.get(spec['key']) if columns else r - 1
            _cell(ws, r, c, value, key=spec['key'], data_type=spec.get('dataType'))
        ws.row_dimensions[r].height = 48
        _highlight(ws, r, source, dataset, end=len(effective))
    _finish(ws)
    return ws


def _reply_level(value):
    if type(value) in (int, float):
        return value >= 2
    if isinstance(value, str):
        return value in ('二级评论', '2级评论', '二级回复') or (
            value.isascii() and value.isdigit() and int(value) >= 2)
    return False


def _relationships(rows):
    index = defaultdict(list)
    for i, row in enumerate(rows):
        note, cid = _id_key(row.get('note_id')), _id_key(row.get('comment_id'))
        if note is not None and cid is not None:
            index[(note, cid)].append(i)
    result = []
    for i, row in enumerate(rows):
        note, cid, root = (_id_key(row.get(k)) for k in ('note_id', 'comment_id', 'thread_root_id'))
        status = row.get('thread_root_status')
        reasons = []
        if status not in (None, '', 'resolved'):
            reasons.append(f'thread_root_status={_text(status)}')
        if note is None or cid is None or root is None:
            reasons.append('缺少有效笔记ID、评论ID或thread_root_id')
        if note is not None and cid is not None and len(index[(note, cid)]) > 1:
            reasons.append('重复评论ID（保留每条筛选记录）')
        parent = _id_key(row.get('parent_comment_id'))
        if root == cid and cid is not None:
            if parent is not None or _reply_level(row.get('comment_level')):
                reasons.append('一级ID与父评/层级冲突')
        # Validate only available exact same-note ID edges. A parent outside the
        # selected records is normal: backend root context is authoritative.
        seen = {cid}
        cursor = parent
        while cursor is not None and not reasons:
            if cursor in seen:
                reasons.append('父链循环')
                break
            seen.add(cursor)
            candidates = index.get((note, cursor), [])
            if not candidates:
                break
            if len(candidates) != 1:
                reasons.append('父链ID存在歧义')
                break
            ancestor = rows[candidates[0]]
            ancestor_root = _id_key(ancestor.get('thread_root_id'))
            if ancestor_root is not None and ancestor_root != root:
                reasons.append('父链与thread_root_id冲突')
                break
            cursor = _id_key(ancestor.get('parent_comment_id'))
        roots = index.get((note, root), []) if note is not None and root is not None else []
        if len(roots) > 1:
            reasons.append('一级评论ID存在歧义')
        elif roots:
            root_row = rows[roots[0]]
            if (_id_key(root_row.get('parent_comment_id')) is not None
                    or _id_key(root_row.get('thread_root_id')) not in (None, root)
                    or root_row.get('thread_root_status') not in (None, '', 'resolved')):
                reasons.append('一级评论关系冲突')
        # The backend uses own-ID as an isolation key for anomalous chains.
        # That key is NOT evidence of an L1 role. Real reply metadata wins.
        if parent is not None or _reply_level(row.get('comment_level')):
            role = '二级回复'
        elif not reasons and root == cid and cid is not None:
            role = '本条即一级'
        elif root is not None and cid is not None and root != cid:
            role = '二级回复'
        else:
            role = '待核验'
        root_author, root_content = row.get('thread_root_author'), row.get('thread_root_content')
        if not reasons and roots:
            root_author = rows[roots[0]].get('author')
            root_content = rows[roots[0]].get('content')
        if not reasons and root == cid:
            root_author, root_content = row.get('author'), row.get('content')
        if root_content in (None, ''):
            parent_label = row.get('thread_root_id')
            if reasons and root == cid and parent is not None:
                parent_label = row.get('parent_comment_id')
            if parent_label in (None, ''):
                parent_label = row.get('parent_comment_id')
            root_content = f'一级评论未入库（父评ID：{_text(parent_label) if parent_label is not None else "缺失"}）'
            if root != cid and status != 'resolved':
                reasons.append('一级评论上下文缺失')
        reason = '；'.join(dict.fromkeys(reasons))
        key = ('isolated', i) if reason else ('thread', note, root)
        result.append({'key': key, 'reason': reason, 'role': role,
                       'author': root_author, 'content': root_content})
    return result


def _groups(indices, relations, rows):
    groups = {}
    for i in indices:
        groups.setdefault(relations[i]['key'], []).append(i)
    for members in groups.values():
        members.sort(key=lambda i: (relations[i]['role'] != '本条即一级',
                                   {'high': 0, 'medium': 1}.get(_control(rows[i])[0], 2),
                                   _text(rows[i].get('comment_id', '')), i))
    return list(groups.values())


def _text_height(value, width):
    """Conservative 11pt wrapped-text estimate, in Excel row-height points.

    Column widths approximate Latin character units. Wide/fullwidth characters
    take two units, combining marks zero; explicit newlines retain empty lines.
    This is a layout estimate, not a change to the stored text or its wrapping.
    """
    if value is None or value == '':
        return 0
    capacity = max(1, int(width - 2))  # Allow for cell padding/borders.
    lines = 0
    for line in _text(value).replace('\r\n', '\n').replace('\r', '\n').split('\n'):
        units = sum(0 if unicodedata.combining(char) else
                    2 if unicodedata.east_asian_width(char) in ('W', 'F') else 1
                    for char in line.expandtabs(4))
        lines += max(1, math.ceil(units / capacity))
    return lines * 15 + 8


def _fit_group_height(ws, start, end):
    """Fit the left merged context against TOTAL group height, not each row.

    Share any shortfall across rows with spare height. Never exceed Excel's
    requested 409pt cap; exceptionally large text remains intact in the cells.
    """
    required = max(_text_height(ws.cell(start, col).value,
                               ws.column_dimensions[get_column_letter(col)].width)
                   for col in range(1, 5))
    dimensions = [ws.row_dimensions[r] for r in range(start, end + 1)]
    deficit = required - sum(d.height for d in dimensions)
    available = [d for d in dimensions if d.height < 409]
    while deficit > 0 and available:
        share = math.ceil(deficit / len(available))
        for dimension in available:
            addition = min(409 - dimension.height, share, deficit)
            dimension.height += addition
            deficit -= addition
        available = [d for d in available if d.height < 409]


def _comments(wb, title, indices, rows, relations, mapping):
    ws = _sheet(wb, title, COMMENT_HEADERS, COMMENT_WIDTHS)
    next_row = 2
    for group_no, members in enumerate(_groups(indices, relations, rows), 1):
        start = next_row
        for offset, i in enumerate(members):
            row, relation = rows[i], relations[i]
            level, _ = _control(row)
            reason = _report(row, 'judgment_reason', mapping)
            if relation['reason']:
                reason = (_text(reason) + '；' if reason not in (None, '') else '') + '待核验：' + relation['reason']
            values = [
                group_no, len(members), relation['author'], relation['content'],
                relation['role'], row.get('author'), row.get('content'),
                _report(row, 'translate_comment', mapping),
                {'high': '需要控制', 'medium': '中等需要控制'}.get(level),
                _report(row, 'main_category', mapping), _report(row, 'subcategory', mapping),
                row.get('comment_level'),
                '是' if _boolean(_report(row, 'highly_similar', mapping)) is True else None,
                '是' if _deleted(row, 'comments') is True else None,
                reason, row.get('note_id'), row.get('post__title'), row.get('post__url'),
                row.get('post__author'), row.get('comment_id'), row.get('published_at'), row.get('like_count'),
            ]
            for col, value in enumerate(values, 1):
                # Only first row owns the left context; merges never discard values.
                if col <= 4 and offset:
                    value = None
                cell = _cell(ws, next_row, col, value,
                             key='id' if col in (16, 20) else '')
                if col <= 4:
                    cell.fill = PatternFill('solid', fgColor=GROUP_COLORS[(group_no - 1) % 2])
                    cell.alignment = Alignment(vertical='center', wrap_text=True)
            role_cell = ws.cell(next_row, 5)
            role_cell.fill = PatternFill('solid', fgColor=HEADER if relation['role'] == '本条即一级' else 'F4F7FA')
            if relation['role'] == '本条即一级':
                role_cell.font = Font(name='Calibri', size=11, color='FFFFFF')
            if level:
                ws.cell(next_row, 9).fill = PatternFill('solid', fgColor='F8D7DA' if level == 'high' else 'FFF3CD')
            category = values[9]
            if isinstance(category, str) and category in CATEGORY_COLORS:
                ws.cell(next_row, 10).fill = PatternFill('solid', fgColor=CATEGORY_COLORS[category])
            link = row.get('post__url')
            if isinstance(link, str) and link.startswith(('https://', 'http://')):
                ws.cell(next_row, 18).hyperlink = link
            _highlight(ws, next_row, row, 'comments', end=22, start=5, status_col=14, preserve=(5, 9, 10))
            body_height = max(_text_height(value, COMMENT_WIDTHS[col - 1])
                              for col, value in enumerate(values[4:], 5))
            ws.row_dimensions[next_row].height = min(409, max(72, body_height))
            next_row += 1
        _fit_group_height(ws, start, next_row - 1)
        if len(members) > 1:
            for col in range(1, 5):
                ws.merge_cells(start_row=start, end_row=next_row - 1, start_column=col, end_column=col)
            ws.row_dimensions.group(start + 1, next_row - 1, outline_level=1, hidden=False)
    _finish(ws)
    return ws


def _summary(wb, dataset, rows, relations, columns, context, mapping):
    ws = _sheet(wb, '统计摘要', ('统计项', '值'), (36, 100))
    counts = Counter(_control(row)[0] for row in rows) if dataset == 'comments' else Counter()
    values = [
        ('数据集', dataset), ('统计口径', '当前筛选同快照全部传入记录；不二次过滤、不去重、不推断、不回写'),
        ('输入记录数', len(rows)), ('主表记录数', len(rows)), ('原始筛选记录数', len(rows)),
        ('一级组数', len({r['key'] for r in relations}) if dataset == 'comments' else '不适用'),
        ('high', counts['high']), ('medium', counts['medium']),
        ('控制级别未知', counts['']), ('待核验记录数', sum(bool(r['reason']) for r in relations)),
        ('已删除记录数', sum(_deleted(r, dataset) is True for r in rows)),
        ('未删除记录数', sum(_deleted(r, dataset) is False for r in rows)),
        ('删除状态未知记录数', sum(_deleted(r, dataset) is None for r in rows)),
        ('明确差评记录数', sum(_boolean(r.get('analysis_is_negative')) is True for r in rows)),
        ('明确非差评记录数', sum(_boolean(r.get('analysis_is_negative')) is False for r in rows)),
        ('差评结论未知记录数', sum(_boolean(r.get('analysis_is_negative')) is None for r in rows)),
        ('颜色说明', '左四列保留分组斑马色；角色/控制/类别格保留专属色；主体删除粉色优先于差评淡绿；同时命中时状态格绿底绿边'),
        ('专项控制来源', '仅control_level/control/report_control明确且一致的high/medium；risk_level/is_negative不参与'),
        ('语义差评来源', '仅analysis_is_negative；旧AI is_negative不参与差评标色'),
    ]
    if dataset == 'notes':
        values.append(('专项子表口径', '帖子数据不套用评论专项控制分类；两个专项子表仅保留帖子列头'))
    else:
        for reason, count in Counter(_control(r)[1] for r in rows if not _control(r)[0]).items():
            values.append(('控制未知原因：' + reason, count))
        for reason, count in Counter(r['reason'] for r in relations if r['reason']).items():
            values.append(('关系待核验：' + reason, count))
        categories = Counter(_text(_report(r, 'main_category', mapping)) for r in rows
                             if _report(r, 'main_category', mapping) not in (None, ''))
        values.extend(('Main Category：' + key, count) for key, count in categories.items())
        values.append(('Main Category缺失', sum(_report(r, 'main_category', mapping) in (None, '') for r in rows)))
        for name, aliases in REPORT_FIELDS.items():
            values.append(('字段映射：' + name, mapping.get(name, ' / '.join(aliases))))
    for key, value in context.items():
        values.append(('context.' + str(key), value))
    if type(context.get('count')) is int:
        values.append(('context.count核对', '一致' if context['count'] == len(rows) else
                       '不一致：按传入记录完整导出，调用方需核验快照与分页'))
    for c in columns:
        values.append(('原始字段：' + c['key'], {'label': c.get('label', c['key']), 'dataType': c.get('dataType')}))
    for r, (label, value) in enumerate(values, 2):
        _cell(ws, r, 1, label)
        _cell(ws, r, 2, value)
        ws.row_dimensions[r].height = 36
    _finish(ws)


def _per_note(wb, dataset, rows, relations, mapping):
    headers = ('笔记ID', '原帖标题', '原帖链接', '筛选记录数', '一级组数', 'high', 'medium',
               '控制级别未知', '已删除记录数', 'Main Category', '待核验记录数')
    if dataset == 'notes':
        headers = ('笔记ID', '原帖标题', '原帖链接', '筛选记录数', '已删除记录数')
    ws = _sheet(wb, '分帖汇总', headers)
    for col, width in ((1, 30), (2, 42), (3, 48)):
        ws.column_dimensions[get_column_letter(col)].width = width
    groups = {}
    for i, row in enumerate(rows):
        key = _id_key(row.get('note_id'))
        groups.setdefault(key if key is not None else ('missing', i), []).append(i)
    for excel_row, members in enumerate(groups.values(), 2):
        row = rows[members[0]]
        prefix = 'post__' if dataset == 'comments' else ''
        values = [row.get('note_id'), row.get(prefix + 'title'), row.get(prefix + 'url'), len(members)]
        deleted = sum(_deleted(rows[i], dataset) is True for i in members)
        if dataset == 'comments':
            counts = Counter(_control(rows[i])[0] for i in members)
            categories = list(dict.fromkeys(_text(_report(rows[i], 'main_category', mapping)) for i in members
                                           if _report(rows[i], 'main_category', mapping) not in (None, '')))
            values += [len({relations[i]['key'] for i in members}), counts['high'], counts['medium'],
                       counts[''], deleted, ' / '.join(categories),
                       sum(bool(relations[i]['reason']) for i in members)]
        else:
            values.append(deleted)
        for col, value in enumerate(values, 1):
            _cell(ws, excel_row, col, value, key='note_id' if col == 1 else '')
        if isinstance(values[2], str) and values[2].startswith(('https://', 'http://')):
            ws.cell(excel_row, 3).hyperlink = values[2]
        ws.row_dimensions[excel_row].height = 48
    _finish(ws)


def build_overview_workbook(dataset, rows, columns, context=None) -> bytes:
    """Return XLSX bytes; raise ValueError on invalid/lossy input, never write files.

    dataset is exactly 'comments' or 'notes'. rows may be a one-shot iterable of
    mappings (including sqlite3.Row, copied into local dicts). columns contains
    {key, label, dataType}. Empty columns use a numbered raw-row placeholder.
    context.field_mapping is an explicit canonical-report-field -> exact-row-key
    map; see module docstring. Other context entries are only audit information.
    """
    if dataset not in ('comments', 'notes'):
        raise ValueError("dataset must be 'comments' or 'notes'")
    rows = [dict(row) for row in rows]
    columns = [dict(column) for column in columns]
    context = dict(context or {})
    if len(rows) + 1 > EXCEL_MAX_ROWS or len(columns) > EXCEL_MAX_COLUMNS:
        raise ValueError('input exceeds Excel row/column limits')
    for column in columns:
        if not isinstance(column.get('key'), str) or not column['key']:
            raise ValueError('each column requires a nonempty string key')
    mapping = context.get('field_mapping', {})
    if not isinstance(mapping, Mapping) or any(k not in REPORT_FIELDS or not isinstance(v, str) or not v
                                             for k, v in mapping.items()):
        raise ValueError('field_mapping must map supported report field names to exact nonempty row keys')
    # Preflight ALL snapshot fields, including fields not selected for the raw
    # sheet. Avoid openpyxl silently slicing strings during cell assignment.
    for i, row in enumerate(rows):
        for key, value in row.items():
            if value is not None:
                _check_text(_text(value), f'rows[{i}].{key}')
    wb = Workbook()
    wb.remove(wb.active)
    relations = _relationships(rows) if dataset == 'comments' else []
    if dataset == 'comments':
        _comments(wb, '筛选评论', range(len(rows)), rows, relations, mapping)
        for title, level in (('需要控制_high', 'high'), ('中等需要控制', 'medium')):
            _comments(wb, title, [i for i, row in enumerate(rows) if _control(row)[0] == level],
                      rows, relations, mapping)
    else:
        _flat(wb, '筛选帖子', rows, columns, dataset)
        _flat(wb, '需要控制_high', [], columns, dataset)
        _flat(wb, '中等需要控制', [], columns, dataset)
    _summary(wb, dataset, rows, relations, columns, context, mapping)
    _per_note(wb, dataset, rows, relations, mapping)
    _flat(wb, '原始筛选数据', rows, columns, dataset)
    output = BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()
