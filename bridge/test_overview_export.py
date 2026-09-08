"""In-memory round-trip tests; run with python -B -m unittest bridge.test_overview_export.

No server imports, database access, filesystem output or integration dependencies.
"""

from copy import deepcopy
from io import BytesIO
import json
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook

try:
    from . import overview_export as export
except ImportError:
    import overview_export as export


def record(cid='root', root='root', note='note-A', **extra):
    row = {
        'note_id': note, 'comment_id': cid, 'thread_root_id': root,
        'thread_root_author': '一级作者', 'thread_root_content': '一级原文',
        'thread_root_status': 'resolved', 'parent_comment_id': '' if cid == root else root,
        'author': '作者-' + cid, 'content': '正文-' + cid,
        'comment_level': 1 if cid == root else 2,
        'post__title': '帖子-' + note, 'post__url': 'https://example.test/' + note,
        'post__author': '帖主', 'published_at': '2026-09-07 10:00:00',
        'like_count': 7, 'analysis_is_negative': '否', 'is_deleted': 0,
    }
    row.update(extra)
    return row


def cols(rows):
    return [{'key': key, 'label': key, 'dataType': 'number' if key == 'like_count' else 'text'}
            for key in dict.fromkeys(key for row in rows for key in row)]


class OverviewExportTests(unittest.TestCase):
    def build(self, rows, dataset='comments', columns=None, context=None):
        data = export.build_overview_workbook(dataset, rows, cols(rows) if columns is None else columns, context)
        self.assertIsInstance(data, bytes)
        wb = load_workbook(BytesIO(data), data_only=False)
        self.addCleanup(wb.close)
        return wb

    def stats(self, wb):
        return dict(wb['统计摘要'].iter_rows(min_row=2, values_only=True))

    def main_by_id(self, wb):
        return {row[19].value: row for row in wb['筛选评论'].iter_rows(min_row=2)}

    def test_three_record_group_merges_four_columns_and_outline(self):
        rows = [record('reply-z', control='medium'), record(), record('reply-a', control='high')]
        wb = self.build(rows)
        ws = wb['筛选评论']
        self.assertEqual(tuple(c.value for c in ws[1]), export.COMMENT_HEADERS)
        self.assertEqual(ws.max_column, 22)
        self.assertEqual(ws.max_row, 4)
        self.assertEqual({str(r) for r in ws.merged_cells.ranges}, {'A2:A4', 'B2:B4', 'C2:C4', 'D2:D4'})
        self.assertEqual(ws['B2'].value, 3)
        self.assertEqual(ws['C2'].value, '作者-root')
        self.assertEqual(ws['D2'].value, '正文-root')
        self.assertEqual([ws.cell(i, 20).value for i in range(2, 5)], ['root', 'reply-a', 'reply-z'])
        self.assertEqual([ws.cell(i, 5).value for i in range(2, 5)], ['本条即一级', '二级回复', '二级回复'])
        self.assertFalse(ws.sheet_properties.outlinePr.summaryBelow)
        self.assertEqual(ws.row_dimensions[2].outlineLevel, 0)
        for r in (3, 4):
            self.assertEqual(ws.row_dimensions[r].outlineLevel, 1)
            self.assertFalse(ws.row_dimensions[r].hidden)
        self.assertEqual(self.stats(wb)['一级组数'], 1)

    def test_same_root_across_notes_never_mix(self):
        rows = [record('a', note='N1'), record('b', note='N2'), record('c', note='N1')]
        wb = self.build(rows)
        ws = wb['筛选评论']
        self.assertEqual(self.stats(wb)['一级组数'], 2)
        self.assertEqual({str(r) for r in ws.merged_cells.ranges}, {'A2:A3', 'B2:B3', 'C2:C3', 'D2:D3'})
        self.assertEqual([ws.cell(i, 16).value for i in range(2, 5)], ['N1', 'N1', 'N2'])
        self.assertEqual(wb['分帖汇总']['D2'].value, 2)
        self.assertEqual(wb['分帖汇总']['D3'].value, 1)

    def test_missing_root_isolated_even_when_content_identical(self):
        rows = [record('a', root=None, thread_root_status='missing'),
                record('b', root=None, thread_root_status='missing')]
        wb = self.build(rows)
        self.assertEqual(len(wb['筛选评论'].merged_cells.ranges), 0)
        self.assertEqual(self.stats(wb)['待核验记录数'], 2)
        self.assertEqual(self.stats(wb)['一级组数'], 2)
        for r in (2, 3):
            self.assertIn('待核验', wb['筛选评论'].cell(r, 15).value)

    def test_missing_parent_context_placeholder_no_synthetic_record(self):
        rows = [record('reply', root='comment-000001', thread_root_content='',
                       thread_root_author='', thread_root_status='missing')]
        wb = self.build(rows)
        self.assertEqual(wb['筛选评论'].max_row, 2)
        self.assertEqual(wb['筛选评论']['D2'].value, '一级评论未入库（父评ID：comment-000001）')
        self.assertEqual(self.stats(wb)['待核验记录数'], 1)

    def test_resolved_unselected_parent_used_as_context_only(self):
        rows = [record('child-a', control='high'), record('child-b', control='high')]
        wb = self.build(rows)
        for title in ('筛选评论', '需要控制_high'):
            ws = wb[title]
            self.assertEqual(ws.max_row, 3)
            self.assertEqual(ws['C2'].value, '一级作者')
            self.assertEqual(ws['D2'].value, '一级原文')
            self.assertEqual(ws['E2'].value, '二级回复')
            self.assertEqual(len(ws.merged_cells.ranges), 4)
        self.assertEqual(self.stats(wb)['待核验记录数'], 0)

    def test_all_records_kept_deleted_ordinary_unknown_and_duplicates(self):
        rows = [record(), record('gone', control='high', is_deleted=1),
                record('ordinary'), record('medium', report_control='medium'),
                record('unknown', risk_level='high', is_negative=1), record('ordinary')]
        wb = self.build(rows)
        self.assertEqual(wb['筛选评论'].max_row - 1, len(rows))
        self.assertEqual(wb['原始筛选数据'].max_row - 1, len(rows))
        self.assertEqual(wb['需要控制_high'].max_row - 1, 1)
        self.assertEqual(wb['中等需要控制'].max_row - 1, 1)
        stats = self.stats(wb)
        self.assertEqual(stats['主表记录数'], 6)
        self.assertEqual(stats['控制级别未知'], 4)
        self.assertEqual(stats['已删除记录数'], 1)
        self.assertEqual(stats['待核验记录数'], 2)
        unknown = self.main_by_id(wb)['unknown']
        self.assertIsNone(unknown[8].value)
        self.assertNotEqual(unknown[6].fill.fgColor.rgb, '00' + export.NEGATIVE_COLOR)

    def test_control_fields_exact_conflicts_unknown_and_legacy_not_used(self):
        rows = [record('a', control_level='high'), record('b', control='medium'),
                record('c', report_control='high'), record('d', risk_level='high', is_negative=True),
                record('e', control_level='high', control='medium'),
                record('f', control='HIGH'), record('g', control='需要控制'),
                record('h', control_level='', report_control='medium')]
        wb = self.build(rows)
        stats = self.stats(wb)
        self.assertEqual((stats['high'], stats['medium'], stats['控制级别未知']), (2, 2, 4))
        self.assertEqual(stats['控制未知原因：专项控制字段冲突'], 1)
        self.assertEqual(stats['控制未知原因：缺少明确专项控制字段'], 1)
        for cid in ('d', 'e', 'f', 'g'):
            self.assertIsNone(self.main_by_id(wb)[cid][8].value)

    def test_exact_ids_prefix_whitespace_and_large_numeric(self):
        ids = ['000123456789012345678901', 'comment-000123456789012345678901', ' root ', 'root', '=1+1']
        rows = [record(cid, root=cid) for cid in ids]
        rows.append(record('placeholder', comment_id=123456789012345678901,
                           thread_root_id=123456789012345678901, parent_comment_id='', comment_level=1))
        wb = self.build(rows)
        self.assertEqual(self.stats(wb)['一级组数'], 6)
        actual = list(self.main_by_id(wb))
        self.assertCountEqual(actual, ids + ['123456789012345678901'])
        for row in wb['筛选评论'].iter_rows(min_row=2):
            self.assertEqual(row[19].data_type, 's')
            self.assertEqual(row[19].number_format, '@')

    def test_integer_and_string_ids_are_not_aliased(self):
        rows = [record('x', comment_id=123, thread_root_id=123, comment_level=1, parent_comment_id=''),
                record('123', root='123')]
        wb = self.build(rows)
        self.assertEqual(self.stats(wb)['一级组数'], 2)

    def test_formula_safety_all_sheets_headers_context_nested_raw_and_links(self):
        formulas = ['=HYPERLINK("https://example.test","x")', '+SUM(1,1)', '-1+2', '@SUM(A1)', '=cmd|evil']
        rows = [record(str(i), content=f, post__title=f, translate_comment=f, judgment_reason=f,
                       extra={'formula': f}, post__url='=HYPERLINK("https://example.test")')
                for i, f in enumerate(formulas)]
        columns = cols(rows)
        columns[0]['label'] = '=1+1'
        data = export.build_overview_workbook('comments', rows, columns, {'search': '=1+2', 'snapshot': {'x': '=2+2'}})
        wb = load_workbook(BytesIO(data))
        self.addCleanup(wb.close)
        for ws in wb:
            for row in ws:
                for cell in row:
                    self.assertNotEqual(cell.data_type, 'f', (ws.title, cell.coordinate))
                    if isinstance(cell.value, str):
                        self.assertEqual(cell.data_type, 's')
        self.assertEqual([self.main_by_id(wb)[str(i)][6].value for i in range(5)], formulas)
        self.assertIsNone(wb['筛选评论']['R2'].hyperlink)
        with ZipFile(BytesIO(data)) as z:
            for path in z.namelist():
                if path.startswith('xl/worksheets/sheet') and path.endswith('.xml'):
                    self.assertNotIn(b'<f>', z.read(path))

    def test_required_five_sheets_plus_raw_and_reconciled_statistics(self):
        rows = [record('a', control='high', main_category='Sales pitch'),
                record('b', control='medium', main_category='Allergy'), record('c', note='note-B')]
        wb = self.build(rows, context={'count': 3, 'filter': {'is_deleted': 1}, 'search': 'unused', 'snapshot': 's1'})
        self.assertEqual(wb.sheetnames, ['筛选评论', '需要控制_high', '中等需要控制', '统计摘要', '分帖汇总', '原始筛选数据'])
        stats = self.stats(wb)
        self.assertEqual(stats['context.count核对'], '一致')
        self.assertEqual(stats['context.snapshot'], 's1')
        self.assertEqual(json.loads(stats['context.filter']), {'is_deleted': 1})
        self.assertEqual(stats['Main Category：Allergy'], 1)
        self.assertEqual(stats['Main Category缺失'], 1)
        sums = [sum(wb['分帖汇总'].cell(r, c).value for r in (2, 3)) for c in (4, 5, 6, 7, 8)]
        self.assertEqual(sums, [3, 2, 1, 1, 1])
        for ws in wb:
            self.assertEqual(ws.freeze_panes, 'A2')
            self.assertEqual(ws.auto_filter.ref, ws.dimensions)
            self.assertEqual(ws['A1'].fill.fgColor.rgb, '00' + export.HEADER)
            self.assertTrue(ws['A1'].font.bold)
            self.assertEqual(ws['A1'].font.color.rgb, '00FFFFFF')

    def test_styles_widths_category_colors_and_hyperlinks(self):
        rows = [record('a', root='a', control='high', main_category='Sales pitch'),
                record('b', root='b', control='medium', main_category='Price discrepancy'),
                record('c', root='c', main_category='Allergy')]
        wb = self.build(rows)
        ws = wb['筛选评论']
        self.assertEqual(ws.column_dimensions['D'].width, 46)
        self.assertEqual(ws.column_dimensions['H'].width, 52)
        for r, category in enumerate(export.CATEGORY_COLORS, 2):
            self.assertEqual(ws.cell(r, 10).fill.fgColor.rgb, '00' + export.CATEGORY_COLORS[category])
            self.assertEqual(ws.cell(r, 4).alignment.vertical, 'center')
            self.assertEqual(ws.cell(r, 7).alignment.vertical, 'top')
            self.assertTrue(ws.cell(r, 7).alignment.wrap_text)
            self.assertEqual(ws.cell(r, 7).border.left.style, 'thin')
            self.assertEqual(ws.cell(r, 18).hyperlink.target, rows[r - 2]['post__url'])
            self.assertEqual(wb['分帖汇总']['C2'].hyperlink.target, rows[0]['post__url'])
        self.assertEqual(ws['A2'].fill.fgColor.rgb, '00E7EEF4')
        self.assertEqual(ws['A3'].fill.fgColor.rgb, '00DCE6EF')
        self.assertEqual(ws['I3'].fill.fgColor.rgb, '00FFF3CD')

    def test_deleted_pink_priority_negative_green_marker_left_unchanged(self):
        rows = [record('both', root='both', is_deleted='1', analysis_is_negative='是'),
                record('neg', root='neg', is_deleted='0', analysis_is_negative='是'),
                record('ordinary', root='ordinary', is_deleted='0', analysis_is_negative='0')]
        wb = self.build(rows)
        byid = self.main_by_id(wb)
        self.assertEqual(byid['both'][6].fill.fgColor.rgb, '00' + export.DELETED_COLOR)
        self.assertEqual(byid['both'][13].fill.fgColor.rgb, '00' + export.NEGATIVE_COLOR)
        self.assertEqual(byid['both'][13].border.left.color.rgb, '00548235')
        self.assertEqual(byid['both'][0].fill.fgColor.rgb, '00E7EEF4')
        self.assertEqual(byid['neg'][6].fill.fgColor.rgb, '00' + export.NEGATIVE_COLOR)
        self.assertEqual(byid['ordinary'][6].fill.patternType, None)
        self.assertIsNone(byid['ordinary'][13].value)

    def test_boolean_zero_false_missing_and_deleted_post(self):
        values = [0, '0', False, 'false', '', None, 'unknown', 1, '1', True]
        rows = [record(str(i), root=str(i), is_deleted=value, analysis_is_negative=value, highly_similar=value)
                for i, value in enumerate(values)]
        rows.append(record('post-deleted', root='post-deleted', post__is_deleted='1'))
        rows.append(record('post-status', root='post-status', post__post_status='已删除'))
        wb = self.build(rows)
        stats = self.stats(wb)
        self.assertEqual(stats['已删除记录数'], 5)
        self.assertEqual(stats['删除状态未知记录数'], 3)
        self.assertEqual(stats['明确差评记录数'], 3)
        for i in range(7):
            row = self.main_by_id(wb)[str(i)]
            self.assertIsNone(row[12].value)
            self.assertIsNone(row[13].value)

    def test_explicit_optional_mapping_no_translation_category_inference(self):
        rows = [record('a', negative_type='过敏', negative_subtype='刺痛',
                       content='过敏 3800 1680 spa high', is_negative=True, risk_level='P0'),
                record('b', english_saved='Existing translation', saved_main='Allergy', saved_sub='Actual allergic reaction',
                       saved_reason='Existing reason', saved_similar='是')]
        mapping = {'translate_comment': 'english_saved', 'main_category': 'saved_main',
                   'subcategory': 'saved_sub', 'judgment_reason': 'saved_reason', 'highly_similar': 'saved_similar'}
        wb = self.build(rows, context={'field_mapping': mapping})
        a, b = self.main_by_id(wb)['a'], self.main_by_id(wb)['b']
        for index in (7, 8, 9, 10, 12, 14):
            self.assertIsNone(a[index].value)
        self.assertEqual([b[i].value for i in (7, 9, 10, 12, 14)],
                         ['Existing translation', 'Allergy', 'Actual allergic reaction', '是', 'Existing reason'])

    def test_backend_relationship_statuses_are_preserved_and_flagged(self):
        rows = [record(str(i), thread_root_status=status) for i, status in
                enumerate(('missing', 'cycle', 'cross_note', 'ambiguous'))]
        wb = self.build(rows)
        self.assertEqual(self.stats(wb)['待核验记录数'], 4)
        self.assertEqual(len(wb['筛选评论'].merged_cells.ranges), 0)
        for r, status in enumerate(('missing', 'cycle', 'cross_note', 'ambiguous'), 2):
            self.assertIn('thread_root_status=' + status, wb['筛选评论'].cell(r, 15).value)

    def test_anomalous_own_id_isolation_key_does_not_claim_level_one(self):
        rows = [record(status, root=status, thread_root_status=status,
                       comment_level=2, parent_comment_id='missing-parent',
                       thread_root_content='一级评论未入库（父评ID：missing-parent）')
                for status in ('missing', 'cycle', 'cross_note', 'ambiguous')]
        rows += [record('level-only', root='level-only', comment_level=3, parent_comment_id='',
                        thread_root_status='missing'),
                 record('parent-only', root='parent-only', comment_level=None, parent_comment_id='p',
                        thread_root_status='cross_note')]
        wb = self.build(rows)
        self.assertEqual(self.stats(wb)['待核验记录数'], 6)
        for r in range(2, 8):
            self.assertEqual(wb['筛选评论'].cell(r, 5).value, '二级回复')
            self.assertIn('待核验', wb['筛选评论'].cell(r, 15).value)
        self.assertEqual(wb['筛选评论']['D2'].value, '一级评论未入库（父评ID：missing-parent）')

    def test_highlight_preserves_role_control_and_category_special_colors(self):
        rows = [record('root', root='root', control='high', main_category='Allergy',
                       is_deleted=1, analysis_is_negative='是'),
                record('reply', root='root', control='medium', main_category='Price discrepancy',
                       is_deleted=1, analysis_is_negative='是')]
        wb = self.build(rows)
        ws = wb['筛选评论']
        self.assertEqual(ws['E2'].fill.fgColor.rgb, '00' + export.HEADER)
        self.assertEqual(ws['E2'].font.color.rgb, '00FFFFFF')
        self.assertEqual(ws['E3'].fill.fgColor.rgb, '00F4F7FA')
        self.assertEqual(ws['I2'].fill.fgColor.rgb, '00F8D7DA')
        self.assertEqual(ws['I3'].fill.fgColor.rgb, '00FFF3CD')
        self.assertEqual(ws['J2'].fill.fgColor.rgb, '00F3D6E8')
        self.assertEqual(ws['J3'].fill.fgColor.rgb, '00FFF3CD')
        for r in (2, 3):
            self.assertEqual(ws.cell(r, 7).fill.fgColor.rgb, '00' + export.DELETED_COLOR)
            self.assertEqual(ws.cell(r, 14).fill.fgColor.rgb, '00' + export.NEGATIVE_COLOR)

    def test_anomalous_own_id_missing_context_uses_real_parent_id(self):
        wb = self.build([record('child', root='child', comment_level=2,
                               parent_comment_id='comment-parent', thread_root_status='missing',
                               thread_root_content=None, thread_root_author=None)])
        self.assertEqual(wb['筛选评论']['D2'].value, '一级评论未入库（父评ID：comment-parent）')
        self.assertEqual(wb['筛选评论']['E2'].value, '二级回复')

    def test_available_parent_cycle_and_conflict_marked_no_text_matching(self):
        rows = [record('a', parent_comment_id='b'), record('b', parent_comment_id='a'),
                record('child', root='other-root', parent_comment_id='parent'),
                record('parent', root='root')]
        wb = self.build(rows)
        byid = self.main_by_id(wb)
        self.assertIn('父链循环', byid['a'][14].value)
        self.assertIn('父链循环', byid['b'][14].value)
        self.assertIn('冲突', byid['child'][14].value)
        self.assertEqual(wb['筛选评论'].max_row, 5)

    def test_parent_with_same_id_other_note_does_not_supply_context(self):
        rows = [record('root', note='other', content='wrong context'),
                record('child', note='wanted', thread_root_content='correct backend context')]
        wb = self.build(rows)
        self.assertEqual(wb['筛选评论']['D3'].value, 'correct backend context')
        self.assertEqual(self.stats(wb)['一级组数'], 2)

    def test_raw_columns_order_duplicate_labels_all_fields_and_missing_values(self):
        row = record('a', payload_json={'nested': ['=1+1', 0, False]}, custom_flag='0')
        row.update({f'extra_{i}': f'value-{i}' for i in range(80)})
        columns = list(reversed(cols([row]))) + [{'key': 'not_present', 'label': 'extra_0', 'dataType': 'text'}]
        wb = self.build([row], columns=columns)
        raw = wb['原始筛选数据']
        self.assertEqual(raw.max_column, len(columns))
        self.assertEqual([c.value for c in raw[1]], [c['label'] for c in columns])
        for col, spec in enumerate(columns, 1):
            expected = row.get(spec['key'])
            actual = raw.cell(2, col).value
            if isinstance(expected, dict):
                self.assertEqual(json.loads(actual), expected)
            elif expected is not None and spec['dataType'] == 'text':
                self.assertEqual(actual, str(expected))
            else:
                self.assertEqual(actual, expected)
        self.assertEqual(wb['筛选评论'].max_column, 22)

    def test_raw_input_order_not_grouped(self):
        rows = [record('b', note='N1'), record('a', note='N2'), record('c', note='N1')]
        wb = self.build(rows, columns=[{'key': 'comment_id', 'label': 'ID', 'dataType': 'text'}])
        self.assertEqual([r[0].value for r in wb['原始筛选数据'].iter_rows(min_row=2)], ['b', 'a', 'c'])
        self.assertEqual([r[19].value for r in wb['筛选评论'].iter_rows(min_row=2)], ['b', 'c', 'a'])

    def test_notes_dataset_uses_columns_without_comment_impersonation(self):
        rows = [{'note_id': '0001234567890123456789', 'title': '=2+2', 'content': '帖子正文',
                 'url': 'https://example.test/note', 'is_deleted': '1', 'analysis_is_negative': '是'},
                {'note_id': 'other', 'title': '普通帖', 'control': 'high'}]
        columns = cols(rows)
        wb = self.build(rows, 'notes', columns=columns)
        self.assertEqual(wb.sheetnames[0], '筛选帖子')
        self.assertNotIn('筛选评论', wb.sheetnames)
        self.assertEqual([c.value for c in wb['筛选帖子'][1]], [c['label'] for c in columns])
        self.assertEqual(wb['筛选帖子'].max_row, 3)
        self.assertEqual(wb['筛选帖子']['A2'].value, rows[0]['note_id'])
        self.assertEqual(wb['筛选帖子']['B2'].data_type, 's')
        self.assertEqual(wb['需要控制_high'].max_row, 1)
        self.assertEqual(wb['中等需要控制'].max_row, 1)
        self.assertEqual(wb['分帖汇总'].max_column, 5)
        self.assertEqual(self.stats(wb)['一级组数'], '不适用')
        self.assertEqual(self.stats(wb)['主表记录数'], 2)

    def test_empty_comments_and_notes_and_zero_columns(self):
        for dataset in ('comments', 'notes'):
            with self.subTest(dataset=dataset):
                wb = self.build([], dataset=dataset)
                self.assertEqual(len(wb.sheetnames), 6)
                self.assertEqual(wb.worksheets[0].max_row, 1)
                self.assertEqual(self.stats(wb)['输入记录数'], 0)
                self.assertEqual(wb['原始筛选数据'].max_row, 1)
                numbered = self.build([{}, {}], dataset=dataset, columns=[])
                self.assertEqual(numbered['原始筛选数据'].max_row, 3)
                self.assertEqual(numbered['原始筛选数据']['A3'].value, 2)

    def test_no_input_mutation_and_one_shot_rows(self):
        rows = [record('b'), record()]
        columns = cols(rows)
        context = {'filter': {'children': [{'key': 'x'}]}, 'count': 999}
        before = deepcopy((rows, columns, context))
        data = export.build_overview_workbook('comments', (r for r in rows), columns, context)
        self.assertEqual((rows, columns, context), before)
        wb = load_workbook(BytesIO(data))
        self.addCleanup(wb.close)
        self.assertIn('不一致', self.stats(wb)['context.count核对'])
        self.assertEqual(self.stats(wb)['主表记录数'], 2)

    def test_long_text_errors_not_truncated_even_unselected_field(self):
        with self.assertRaisesRegex(ValueError, r'rows\[0\].content.*32767'):
            self.build([record(content='x' * 32768)])
        with self.assertRaisesRegex(ValueError, 'unselected.*32767'):
            self.build([record(unselected='x' * 32768)], columns=[{'key': 'comment_id', 'label': 'ID'}])
        with self.assertRaisesRegex(ValueError, '32767'):
            self.build([record()], context={'search': 'x' * 32768})
        with self.assertRaisesRegex(ValueError, '32767'):
            self.build([record(content='😀' * 16384)])
        wb = self.build([record(content='x' * 32767)])
        self.assertEqual(len(wb['筛选评论']['G2'].value), 32767)

    def test_299_chinese_characters_fit_single_comment_and_keep_special_colors(self):
        text = '评' * 299
        wb = self.build([record(content=text, control='high', main_category='Allergy',
                                is_deleted=1, analysis_is_negative='是')])
        for title in ('筛选评论', '需要控制_high'):
            ws = wb[title]
            # D width46: 22 Chinese characters per line => 14 lines, 218pt.
            self.assertEqual(ws.row_dimensions[2].height, 218)
            self.assertEqual(ws['D2'].value, text)
            self.assertEqual(ws['G2'].value, text)
            self.assertEqual(ws['I2'].fill.fgColor.rgb, '00F8D7DA')
            self.assertEqual(ws['J2'].fill.fgColor.rgb, '00F3D6E8')
            self.assertEqual(ws['E2'].font.color.rgb, '00FFFFFF')

    def test_long_parent_short_replies_fit_combined_height_not_each_row(self):
        parent = '父' * 299
        rows = [record(cid, content='短回复', thread_root_content=parent,
                       control='high', main_category='Allergy') for cid in ('a', 'b')]
        wb = self.build(rows)
        for title in ('筛选评论', '需要控制_high'):
            ws = wb[title]
            self.assertEqual(ws.max_row, 3)
            self.assertEqual({str(r) for r in ws.merged_cells.ranges},
                             {'A2:A3', 'B2:B3', 'C2:C3', 'D2:D3'})
            self.assertEqual(ws['D2'].value, parent)
            self.assertEqual(sum(ws.row_dimensions[r].height for r in (2, 3)), 218)
            for r in (2, 3):
                self.assertEqual(ws.row_dimensions[r].height, 109)
                self.assertEqual(ws.cell(r, 7).value, '短回复')
                self.assertEqual(ws.cell(r, 9).fill.fgColor.rgb, '00F8D7DA')
                self.assertEqual(ws.cell(r, 10).fill.fgColor.rgb, '00F3D6E8')
            self.assertEqual(ws.row_dimensions[3].outlineLevel, 1)

    def test_body_fields_newlines_east_asian_width_and_height_bounds(self):
        self.assertGreater(export._text_height('中' * 100, 48), export._text_height('a' * 100, 48))
        self.assertEqual(export._text_height('a\r\n\r\nb', 48), 53)
        self.assertEqual(export._text_height('e\u0301' * 46, 48), export._text_height('e' * 46, 48))
        for field in ('content', 'translate_comment', 'judgment_reason', 'post__title'):
            with self.subTest(field=field):
                wb = self.build([record('reply', **{field: '\n'.join(['正文'] * 12)})])
                self.assertEqual(wb['筛选评论'].row_dimensions[2].height, 188)
        wb = self.build([record('short'), record('long', content='文' * 2000)])
        self.assertEqual(self.main_by_id(wb)['long'][6].value, '文' * 2000)
        for row in wb['筛选评论'].iter_rows(min_row=2):
            self.assertEqual(wb['筛选评论'].row_dimensions[row[0].row].height,
                             409 if row[19].value == 'long' else 72)

    def test_parent_height_spare_capacity_redistribution_and_existing_height(self):
        rows = [record('a', content='文' * 2000, thread_root_content='父' * 800),
                record('b', content='短', thread_root_content='父' * 800)]
        wb = self.build(rows)
        ws = wb['筛选评论']
        self.assertEqual(ws.row_dimensions[2].height, 409)
        self.assertEqual(ws.row_dimensions[3].height, 154)
        self.assertEqual(ws['D2'].value, '父' * 800)
        # Three short rows already provide enough height for a 200-character parent.
        wb = self.build([record(cid, content='短', thread_root_content='父' * 200) for cid in ('a', 'b', 'c')])
        self.assertEqual([wb['筛选评论'].row_dimensions[r].height for r in (2, 3, 4)], [72, 72, 72])

    def test_invalid_xml_and_unsupported_dataset_mapping_and_limits(self):
        with self.assertRaisesRegex(ValueError, 'XML-incompatible'):
            self.build([record(content='abc\x00def')])
        with self.assertRaisesRegex(ValueError, 'dataset'):
            self.build([], dataset='bad')
        with self.assertRaisesRegex(ValueError, 'field_mapping'):
            self.build([], context={'field_mapping': {'control': 'risk_level'}})
        with self.assertRaisesRegex(ValueError, 'key'):
            self.build([], columns=[{'label': 'no key'}])
        with patch.object(export, 'EXCEL_MAX_ROWS', 2):
            with self.assertRaisesRegex(ValueError, 'limit'):
                self.build([record(), record('a')])
        with self.assertRaisesRegex(ValueError, 'non-finite'):
            self.build([record(like_count=float('inf'))])

    def test_many_records_no_loss_no_dedup_independent_of_group_reorder(self):
        rows = [record(str(i), root='r-' + str(i % 13), note='n-' + str(i % 5),
                       control_level=('high', 'medium', 'unknown')[i % 3],
                       is_deleted=i % 2, analysis_is_negative='是' if i % 4 else '否')
                for i in range(130)]
        wb = self.build(rows)
        self.assertEqual(wb['筛选评论'].max_row - 1, 130)
        self.assertEqual(wb['原始筛选数据'].max_row - 1, 130)
        self.assertCountEqual(self.main_by_id(wb), [str(i) for i in range(130)])
        self.assertEqual(self.stats(wb)['一级组数'], 65)
        self.assertEqual(self.stats(wb)['high'] + self.stats(wb)['medium'] + self.stats(wb)['控制级别未知'], 130)
        self.assertEqual(sum(wb['筛选评论'].cell(r, 2).value or 0 for r in range(2, 132)), 130)


if __name__ == '__main__':
    unittest.main()
