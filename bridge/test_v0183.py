import csv
import json
import tempfile
import threading
import time
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from data_relationships import repair_relationship_rows
from server import COMMENT_CSV_HEADERS, NOTE_CSV_HEADERS, MonitorStore, _decode_powershell_output, audit_reply_candidate, classify_reply_context, comment_level_value, replace_with_retry, tag_text, valid_note_id


COMPLETE_EVIDENCE = {
    "allCommentsRequested": True,
    "expandersExhausted": True,
    "scrollExhausted": True,
    "stableRounds": 2,
}


class MediaHandler(BaseHTTPRequestHandler):
    counts = {"image": 0, "video": 0}

    def do_GET(self):
        if self.path.startswith("/post.jpg"):
            self.counts["image"] += 1
            body, content_type = b"post-image", "image/jpeg"
        elif self.path.startswith("/post.mp4"):
            self.counts["video"] += 1
            body, content_type = b"video-data" * 128, "video/mp4"
        else:
            self.send_response(404); self.end_headers(); return
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, *_args):
        return


class FakeAI:
    def complete_json(self, _settings, _messages):
        return {
            "overview": "帖子介绍产品使用体验，评论整体中性。", "sentiment": "neutral",
            "key_points": ["正文核心"], "comment_consensus": ["评论共识"],
            "disagreements": [], "risks": ["暂无明显风险"], "actions": ["持续观察"],
            "representative_comments": ["代表评论"],
        }


class FakeReplyAI:
    def complete_json(self, _settings, _messages):
        return {"need": "想确认价格", "candidates": [
            {"reply": "理解你对价格的顾虑。你发下产品全名，我们按门店公示价帮你核对。", "style": "直答", "why": "接住价格疑问"},
            {"reply": "确实要先把价格问清楚。方便的话发张包装图，我们帮你确认具体款式。", "style": "共情", "why": "先确认产品"},
            {"reply": "价格会跟系列和容量有关。告诉我们产品名和门店，我们再核对当日标价。", "style": "稳妥", "why": "避免编价"}
        ], "risk_notes": []}


class V0183Tests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        self.store = MonitorStore(root / "monitor.sqlite3", root / "exports", ai_client=FakeAI())
        self.server = ThreadingHTTPServer(("127.0.0.1", 0), MediaHandler)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        MediaHandler.counts = {"image": 0, "video": 0}
        self.base = f"http://127.0.0.1:{self.server.server_port}"

    def tearDown(self):
        self.server.shutdown(); self.server.server_close(); self.tmp.cleanup()

    def _configure_csv(self, name: str = "master"):
        notes_path, comments_path = self.store.configure_data_files(Path(self.tmp.name) / f"{name}_笔记总表.csv")
        self.store._ensure_seed_workbook(notes_path)
        return notes_path, comments_path

    def _csv_rows(self, path: Path, headers: list[str]):
        return self.store._read_csv_table(path, headers)[1]

    def test_media_is_idempotent_and_only_repairs_missing_video(self):
        note = {"noteId": "abcdef123456", "title": "视频帖子", "content": "正文",
                "imageUrls": [self.base + "/post.jpg"], "videoUrls": [self.base + "/post.mp4"]}
        first = self.store._download_note_media(note)
        self.assertEqual((1, 1), (first["imageCount"], first["videoCount"]))
        self.store._download_note_media(note)
        self.assertEqual({"image": 1, "video": 1}, MediaHandler.counts)
        Path(first["folder"], "video-01.mp4").unlink()
        repaired = self.store._download_note_media(note)
        self.assertEqual({"image": 1, "video": 2}, MediaHandler.counts)
        self.assertEqual(1, repaired["downloadedCount"])
        self.assertEqual(1, repaired["skippedCount"])

    def test_title_only_scan_prefers_existing_note_id_over_changed_title(self):
        timestamp = "2026-08-21T12:00:00+08:00"
        with self.store._session() as db:
            db.execute("""
                INSERT INTO notes
                (note_id, url, title, first_seen_at, last_seen_at, status, is_relevant, source,
                 title_key, content_key, title_content_key, pull_status)
                VALUES (?, ?, ?, ?, ?, 'known', 1, 'existing_xlsx', ?, '', '', 'synced')
            """, ("existing123456", "https://www.xiaohongshu.com/explore/existing123456",
                  "Excel保存的旧标题", timestamp, timestamp, "excel保存的旧标题"))
        result = self.store.scan({
            "titleOnly": True, "returnAllStatuses": True, "keyword": "samplebrand",
            "notes": [{"noteId": "existing123456", "url": "https://www.xiaohongshu.com/explore/existing123456",
                       "title": "小红书当前展示的新标题", "content": ""}]
        })
        self.assertTrue(result["ok"])
        self.assertEqual(1, len(result["statuses"]))
        self.assertTrue(result["statuses"][0]["inExcel"])
        self.assertEqual("known", result["statuses"][0]["status"])
        self.assertEqual("帖子ID", result["statuses"][0]["matchLabel"])

    def test_title_only_scan_never_overwrites_dom_sourced_pulled_note(self):
        timestamp = "2026-08-21T12:00:00+08:00"
        canonical_payload = {
            "noteId": "dompulled123456",
            "url": "https://www.xiaohongshu.com/search_result/dompulled123456?xsec_token=CANONICAL",
            "title": "总表完整标题", "author": "总表作者", "content": "总表完整正文",
        }
        with self.store._session() as db:
            db.execute("""
                INSERT INTO notes
                (note_id,url,title,author,content,tags,keyword,first_seen_at,last_seen_at,status,
                 is_relevant,source,title_key,content_key,title_content_key,pull_status,payload_json)
                VALUES(?,?,?,?,?,?,?,?,?,'known',1,'dom',?,?,?,'synced',?)
            """, (
                canonical_payload["noteId"], canonical_payload["url"], canonical_payload["title"],
                canonical_payload["author"], canonical_payload["content"], "完整话题", "来源词",
                timestamp, timestamp, "总表完整标题", "总表完整正文", "总表完整标题总表完整正文",
                json.dumps(canonical_payload, ensure_ascii=False),
            ))
        result = self.store.scan({
            "titleOnly": True, "returnAllStatuses": True, "keyword": "samplebrand",
            "notes": [{
                "noteId": canonical_payload["noteId"],
                "url": "https://www.xiaohongshu.com/search_result/dompulled123456?xsec_token=ROTATED",
                "title": "卡片截断标题", "author": "卡片作者", "content": "卡片截断正文",
                "tags": ["卡片话题"],
            }],
        })
        self.assertTrue(result["statuses"][0]["inExcel"])
        with self.store._session() as db:
            stored = dict(db.execute(
                "SELECT url,title,author,content,tags,keyword,payload_json FROM notes WHERE note_id=?",
                (canonical_payload["noteId"],),
            ).fetchone())
        self.assertEqual(canonical_payload["url"], stored["url"])
        self.assertEqual("总表完整标题", stored["title"])
        self.assertEqual("总表作者", stored["author"])
        self.assertEqual("总表完整正文", stored["content"])
        self.assertEqual("完整话题", stored["tags"])
        self.assertEqual("来源词", stored["keyword"])
        self.assertEqual(canonical_payload, json.loads(stored["payload_json"]))

    def test_same_title_different_note_id_is_never_marked_as_pulled(self):
        timestamp = "2026-08-27T15:00:00+08:00"
        original_payload = {
            "noteId": "pulledtitle123", "url": "https://www.xiaohongshu.com/explore/pulledtitle123",
            "title": "samplebrand", "author": "已拉取作者", "content": "已拉取正文",
        }
        with self.store._session() as db:
            db.execute("""
                INSERT INTO notes
                (note_id,url,title,author,content,first_seen_at,last_seen_at,status,is_relevant,source,
                 title_key,content_key,title_content_key,pull_status,payload_json)
                VALUES(?,?,?,?,?,?,?,'known',1,'existing_xlsx','samplebrand','已拉取正文','samplebrand已拉取正文','synced',?)
            """, ("pulledtitle123", original_payload["url"], "samplebrand", "已拉取作者", "已拉取正文",
                  timestamp, timestamp, json.dumps(original_payload, ensure_ascii=False)))

        result = self.store.scan({
            "titleOnly": True, "returnAllStatuses": True, "keyword": "samplebrand",
            "notes": [{"noteId": "deletedtitle456", "url": "https://www.xiaohongshu.com/explore/deletedtitle456",
                       "title": "samplebrand", "author": "另一位作者", "content": ""}]
        })

        status = result["statuses"][0]
        self.assertEqual("deletedtitle456", status["noteId"])
        self.assertEqual("deletedtitle456", status["matchedNoteId"])
        self.assertFalse(status["inExcel"])
        self.assertEqual("not_started", status["pullStatus"])
        self.assertEqual("none", status["matchedBy"])
        with self.store._session() as db:
            stored = json.loads(db.execute(
                "SELECT payload_json FROM notes WHERE note_id='pulledtitle123'"
            ).fetchone()[0])
        self.assertEqual("pulledtitle123", stored["noteId"])
        self.assertEqual("已拉取作者", stored["author"])

    def test_pull_and_csv_sync_never_deduplicate_different_note_ids_by_text(self):
        notes_path, _comments_path = self._configure_csv("strict-identity")
        first = {"noteId": "strictnote123", "url": "https://www.xiaohongshu.com/explore/strictnote123",
                 "title": "完全相同标题", "content": "完全相同正文", "detailRead": True}
        second = {**first, "noteId": "strictnote456",
                  "url": "https://www.xiaohongshu.com/explore/strictnote456"}
        self.store._sync_pull_to_xlsx(first, [], {})
        self.store._sync_pull_to_xlsx(second, [], {})
        rows = self._csv_rows(notes_path, NOTE_CSV_HEADERS)
        self.assertEqual({"strictnote123", "strictnote456"}, {row["笔记ID"] for row in rows})
        self.store.confirm(first)
        resolved, matched_by = self.store._resolve_pull_identity(second)
        self.assertEqual("strictnote456", resolved)
        self.assertEqual("new", matched_by)

    def test_placeholder_note_ids_are_rejected(self):
        for value in ("undefined", "null", "None", "UNKNOWN", "nan"):
            self.assertEqual("", valid_note_id(value))

    def test_startup_repairs_and_archives_cross_note_payload_identity(self):
        timestamp = "2026-08-27T15:00:00+08:00"
        with self.store._session() as db:
            db.execute("""
                INSERT INTO notes(note_id,url,title,author,content,first_seen_at,last_seen_at,payload_json)
                VALUES(?,?,?,?,?,?,?,?)
            """, ("payloadnote123", "https://www.xiaohongshu.com/explore/payloadnote123",
                  "正确标题", "正确作者", "正确正文", timestamp, timestamp,
                  json.dumps({"noteId": "othernote456", "url": "https://www.xiaohongshu.com/explore/othernote456",
                              "title": "错误标题"}, ensure_ascii=False)))
        repaired_store = MonitorStore(self.store.db_path, self.store.export_dir, ai_client=FakeAI())
        with repaired_store._session() as db:
            payload = json.loads(db.execute(
                "SELECT payload_json FROM notes WHERE note_id='payloadnote123'"
            ).fetchone()[0])
            archived = db.execute(
                "SELECT COUNT(*) FROM data_repair_archive WHERE repair_id LIKE 'v0.25.1-note-payload-identity%'"
            ).fetchone()[0]
        self.assertEqual("payloadnote123", payload["noteId"])
        self.assertIn("/payloadnote123", payload["url"])
        self.assertEqual("正确标题", payload["title"])
        self.assertGreaterEqual(archived, 1)

    def test_pulled_note_status_hydrates_artifacts_and_comments(self):
        root = self.store.export_dir.parent
        media_dir = root / "materials" / "pulled123456"
        media_dir.mkdir(parents=True)
        (media_dir / "image-01.jpg").write_bytes(b"image")
        excel_path = root / "master.xlsx"
        timestamp = "2026-08-24T12:00:00+08:00"
        with self.store._session() as db:
            db.execute("""
                INSERT INTO notes
                (note_id, url, title, author, content, first_seen_at, last_seen_at, status,
                 is_relevant, source, title_key, content_key, title_content_key, pull_status,
                 media_dir, excel_sync_path, payload_json)
                VALUES (?, ?, ?, ?, ?, ?, ?, 'known', 1, 'existing_xlsx', ?, ?, '', 'synced', ?, ?, ?)
            """, ("pulled123456", "https://www.xiaohongshu.com/explore/pulled123456",
                  "已拉取帖子", "作者", "正文", timestamp, timestamp, "已拉取帖子", "正文",
                  str(media_dir), str(excel_path), json.dumps({"likeCount": 8}, ensure_ascii=False)))
            db.execute("""
                INSERT INTO comments
                (comment_id, note_id, content, author, first_seen_at, last_seen_at, payload_json)
                VALUES (?, ?, ?, ?, ?, ?, ?)
            """, ("comment-pulled-1", "pulled123456", "本地评论", "用户", timestamp, timestamp,
                  json.dumps({"commentId": "comment-pulled-1", "content": "本地评论"}, ensure_ascii=False)))
        result = self.store.note_status("pulled123456")
        self.assertTrue(result["inExcel"])
        self.assertEqual(str(media_dir), result["mediaDir"])
        self.assertEqual(["image-01.jpg"], result["mediaFiles"])
        self.assertEqual(str(excel_path), result["excelPath"])
        self.assertEqual("pulled123456", result["note"]["noteId"])
        self.assertEqual(1, result["commentCount"])
        self.assertEqual("comment-pulled-1", result["commentRows"][0]["commentId"])

    def _seed_pulled_note_with_comments(self):
        note = {
            "noteId": "audit123456", "url": "https://www.xiaohongshu.com/explore/audit123456",
            "title": "评论变化测试", "content": "正文", "author": "博主", "detailRead": True,
        }
        self.store.confirm(note)
        old = {"commentId": "comment-old", "author": "旧用户", "content": "已经消失", "publishedAt": "08-20"}
        kept = {"commentId": "comment-kept", "author": "保留用户", "content": "仍然存在", "publishedAt": "08-21"}
        self.store.upsert_comments({
            "noteId": note["noteId"], "comments": [old, kept],
            "expectedCount": 2, "status": "likely_complete"
        })
        return note, old, kept

    def test_comment_compare_reports_new_and_confirmed_removed(self):
        note, old, kept = self._seed_pulled_note_with_comments()
        fresh = {"commentId": "comment-new", "author": "新用户", "content": "新增评论", "publishedAt": "08-24"}
        result = self.store.compare_comments({
            "noteId": note["noteId"], "comments": [kept, fresh],
            "expectedCount": 2, "collectionEvidence": COMPLETE_EVIDENCE,
            "status": "likely_complete"
        })
        self.assertTrue(result["hasChanges"])
        self.assertTrue(result["canPrune"])
        self.assertEqual(1, result["newCount"])
        self.assertEqual(1, result["removedCount"])
        self.assertEqual("comment-old", result["removedComments"][0]["commentId"])

    def test_different_comment_ids_never_merge_even_when_text_matches(self):
        note, old, kept = self._seed_pulled_note_with_comments()
        _notes_path, comments_path = self._configure_csv("strict-comment-id")
        self.store._sync_pull_to_xlsx(note, [old, kept], {"folder": "", "files": []})
        replacement = {**old, "commentId": "comment-new-identity"}
        compared = self.store.compare_comments({
            "noteId": note["noteId"], "comments": [replacement, kept],
            "expectedCount": 2, "collectionEvidence": COMPLETE_EVIDENCE,
            "status": "likely_complete",
        })
        self.assertEqual((1, 1), (compared["newCount"], compared["removedCount"]))
        self.assertEqual("comment-old", compared["removedComments"][0]["commentId"])
        self.store.sync_comment_snapshot({
            "noteId": note["noteId"], "note": note, "comments": [replacement, kept],
            "expectedCount": 2, "collectionEvidence": COMPLETE_EVIDENCE,
            "status": "likely_complete",
        })
        rows = {row["笔记评论ID"]: row for row in self._csv_rows(comments_path, COMMENT_CSV_HEADERS)}
        self.assertEqual({"comment-old", "comment-kept", "comment-new-identity"}, set(rows))
        self.assertEqual("已删除", rows["comment-old"]["评论状态"])
        self.assertEqual("存在", rows["comment-new-identity"]["评论状态"])

    def test_partial_comment_compare_never_confirms_deletion(self):
        note, _old, kept = self._seed_pulled_note_with_comments()
        result = self.store.compare_comments({
            "noteId": note["noteId"], "comments": [kept],
            "expectedCount": 2, "status": "partial"
        })
        self.assertFalse(result["canPrune"])
        self.assertEqual(0, result["removedCount"])
        self.assertEqual(1, result["pendingRemovedCount"])
        self.assertFalse(result["hasChanges"])

    def test_post_metadata_change_does_not_mark_comments_changed(self):
        note, old, kept = self._seed_pulled_note_with_comments()
        with self.store._session() as db:
            db.execute("UPDATE notes SET payload_json=? WHERE note_id=?", (
                json.dumps({**note, "likeCount": 1}, ensure_ascii=False), note["noteId"]
            ))
        result = self.store.compare_comments({
            "noteId": note["noteId"], "note": {**note, "likeCount": 2},
            "comments": [old, kept], "expectedCount": 2, "status": "likely_complete"
        })
        self.assertTrue(result["noteChanged"])
        self.assertTrue(result["hasChanges"])
        self.assertFalse(result["commentHasChanges"])
        self.assertEqual((0, 0, 0), (result["newCount"], result["removedCount"], result["changedCount"]))

    def test_ignored_pulled_note_cascades_and_selectively_restores_presence(self):
        note, old, kept = self._seed_pulled_note_with_comments()
        notes_path, comments_path = self._configure_csv("ignored-presence")
        self.store._sync_pull_to_xlsx(note, [old, kept], {"folder": "", "files": []})
        with self.store._session() as db:
            db.execute("UPDATE notes SET source='existing_xlsx',pull_status='synced',access_status='unreachable' WHERE note_id=?", (note["noteId"],))
        ignored = self.store.ignore({"noteId": note["noteId"]})
        self.assertTrue(ignored["consistencyVerified"])
        self.assertEqual(2, ignored["commentsMarkedDeleted"])
        self.assertEqual([], self.store.list_unreachable_notes())
        note_row = self._csv_rows(notes_path, NOTE_CSV_HEADERS)[0]
        self.assertEqual("已删除", note_row["帖子状态"])
        self.assertEqual({"已删除"}, {row["评论状态"] for row in self._csv_rows(comments_path, COMMENT_CSV_HEADERS)})
        with self.store._session() as db:
            self.assertEqual(2, db.execute(
                "SELECT COUNT(*) FROM comments WHERE note_id=? AND is_deleted=1", (note["noteId"],)
            ).fetchone()[0])
        restored = self.store.restore({"noteId": note["noteId"]})
        self.assertTrue(restored["consistencyVerified"])
        self.assertEqual(2, restored["commentsRestored"])
        with self.store._session() as db:
            status = db.execute("SELECT status FROM notes WHERE note_id=?", (note["noteId"],)).fetchone()[0]
        self.assertEqual("known", status)
        self.assertEqual({"存在"}, {row["评论状态"] for row in self._csv_rows(comments_path, COMMENT_CSV_HEADERS)})
        self.assertEqual(note["noteId"], self.store.list_unreachable_notes()[0]["note_id"])

    def test_restore_ignored_note_does_not_revive_previously_missing_comment(self):
        note, old, kept = self._seed_pulled_note_with_comments()
        _notes_path, comments_path = self._configure_csv("ignored-selective-restore")
        self.store._sync_pull_to_xlsx(note, [old, kept], {"folder": "", "files": []})
        self.store.sync_comment_snapshot({
            "noteId": note["noteId"], "note": note, "comments": [kept],
            "expectedCount": 1, "collectionEvidence": COMPLETE_EVIDENCE,
            "status": "likely_complete",
        })
        with self.store._session() as db:
            db.execute("UPDATE notes SET source='existing_xlsx',pull_status='synced' WHERE note_id=?", (note["noteId"],))
        self.store.ignore({"noteId": note["noteId"]})
        restored = self.store.restore({"noteId": note["noteId"]})
        self.assertEqual(1, restored["commentsRestored"])
        rows = {row["笔记评论ID"]: row for row in self._csv_rows(comments_path, COMMENT_CSV_HEADERS)}
        self.assertEqual("已删除", rows[old["commentId"]]["评论状态"])
        self.assertEqual("存在", rows[kept["commentId"]]["评论状态"])

    def test_comment_sync_marks_deleted_without_losing_history_or_semantic_fields(self):
        note, old, kept = self._seed_pulled_note_with_comments()
        _notes_path, comments_path = self._configure_csv("comments")
        media_dir = Path(self.tmp.name) / "posts_materials" / "评论状态"
        media_dir.mkdir(parents=True)
        with self.store._session() as db:
            db.execute("UPDATE notes SET media_dir=? WHERE note_id=?", (str(media_dir), note["noteId"]))
        self.store._sync_pull_to_xlsx(note, [old, kept], {"folder": str(media_dir), "files": []})
        headers, rows = self.store._read_csv_table(comments_path, COMMENT_CSV_HEADERS)
        old_row = next(row for row in rows if row["笔记评论ID"] == "comment-old")
        old_row.update({"语义分析次数": "2", "分析结论是否差评": "是",
                        "差评类型": "Sales pitch", "差评子类型": "Aggressive pull-in"})
        self.store._replace_csv_table(comments_path, headers, rows, "semantic-test")
        with self.store._session() as db:
            db.execute("""UPDATE comments SET semantic_analysis_count=2,analysis_is_negative='是',
                       negative_type='Sales pitch',negative_subtype='Aggressive pull-in'
                       WHERE comment_id='comment-old'""")
        fresh = {"commentId": "comment-new", "author": "新用户", "content": "新增评论", "publishedAt": "08-24"}
        result = self.store.sync_comment_snapshot({
            "noteId": note["noteId"], "note": note, "comments": [kept, fresh],
            "expectedCount": 2, "collectionEvidence": COMPLETE_EVIDENCE,
            "status": "likely_complete"
        })
        self.assertEqual("latest", result["status"])
        self.assertEqual("not_started", result["pullStatus"])
        self.assertEqual((1, 1), (result["newCount"], result["removedCount"]))
        db_rows = {row["comment_id"]: row for row in self.store.list_comments(note["noteId"])}
        self.assertEqual({"comment-old", "comment-kept", "comment-new"}, set(db_rows))
        self.assertEqual(1, db_rows["comment-old"]["is_deleted"])
        self.assertEqual("已删除", db_rows["comment-old"]["comment_status"])
        self.assertEqual(2, db_rows["comment-old"]["semantic_analysis_count"])
        csv_rows = {row["笔记评论ID"]: row for row in self._csv_rows(comments_path, COMMENT_CSV_HEADERS)}
        self.assertEqual(set(db_rows), set(csv_rows))
        self.assertEqual("已删除", csv_rows["comment-old"]["评论状态"])
        self.assertEqual("存在", csv_rows["comment-kept"]["评论状态"])
        self.assertEqual("Sales pitch", csv_rows["comment-old"]["差评类型"])
        material_rows = {row["commentId"]: row for row in json.loads(
            (media_dir / "comments.json").read_text(encoding="utf-8")
        )}
        self.assertTrue(material_rows["comment-old"]["isDeleted"])
        self.assertEqual("已删除", material_rows["comment-old"]["commentStatus"])
        self.assertEqual("Sales pitch", material_rows["comment-old"]["negativeType"])
        repeat = self.store.compare_comments({
            "noteId": note["noteId"], "comments": [kept, fresh],
            "expectedCount": 2, "status": "likely_complete"
        })
        self.assertEqual(0, repeat["removedCount"])

        restored = self.store.sync_comment_snapshot({
            "noteId": note["noteId"], "note": note, "comments": [old, kept, fresh],
            "expectedCount": 3, "collectionEvidence": COMPLETE_EVIDENCE,
            "status": "likely_complete"
        })
        self.assertEqual(1, restored["newCount"])
        restored_row = next(row for row in self._csv_rows(comments_path, COMMENT_CSV_HEADERS)
                            if row["笔记评论ID"] == "comment-old")
        self.assertEqual("存在", restored_row["评论状态"])
        self.assertEqual("Sales pitch", restored_row["差评类型"])
        restored_material = {row["commentId"]: row for row in json.loads(
            (media_dir / "comments.json").read_text(encoding="utf-8")
        )}
        self.assertFalse(restored_material["comment-old"]["isDeleted"])

    def test_single_pull_complete_snapshot_marks_missing_comment_deleted(self):
        note, old, kept = self._seed_pulled_note_with_comments()
        _notes_path, comments_path = self._configure_csv("single-complete-presence")
        self.store._sync_pull_to_xlsx(note, [old, kept], {"folder": "", "files": []})
        result = self.store.pull_to_excel({
            "note": note, "comments": [kept], "expectedCount": 1,
            "collectionEvidence": COMPLETE_EVIDENCE,
            "commentStatus": "likely_complete",
        })
        self.assertEqual(1, result["commentMarkedDeleted"])
        self.assertTrue(result["consistencyVerified"])
        self.assertTrue(result["consistency"]["stateHash"])
        csv_rows = {row["笔记评论ID"]: row for row in self._csv_rows(comments_path, COMMENT_CSV_HEADERS)}
        self.assertEqual("已删除", csv_rows["comment-old"]["评论状态"])
        with self.store._session() as db:
            status = db.execute(
                "SELECT comment_status,is_deleted FROM comments WHERE comment_id='comment-old'"
            ).fetchone()
        self.assertEqual(("已删除", 1), tuple(status))

    def test_failed_field_verification_rolls_back_csv_and_sqlite(self):
        note, old, kept = self._seed_pulled_note_with_comments()
        notes_path, comments_path = self._configure_csv("sync-rollback")
        self.store._sync_pull_to_xlsx(note, [old, kept], {"folder": "", "files": []})
        before_files = (notes_path.read_bytes(), comments_path.read_bytes())
        with self.store._session() as db:
            before_db = [tuple(row) for row in db.execute(
                "SELECT comment_id,content,comment_status,is_deleted FROM comments WHERE note_id=? ORDER BY comment_id",
                (note["noteId"],),
            ).fetchall()]
        original_verify = self.store._verify_note_store_consistency
        self.store._verify_note_store_consistency = lambda *_args, **_kwargs: (_ for _ in ()).throw(
            ValueError("forced field mismatch")
        )
        try:
            with self.assertRaisesRegex(ValueError, "forced field mismatch"):
                self.store.sync_comment_snapshot({
                    "noteId": note["noteId"], "note": note,
                    "comments": [kept, {"commentId": "comment-new", "author": "新", "content": "新增"}],
                    "expectedCount": 2, "status": "likely_complete",
                })
        finally:
            self.store._verify_note_store_consistency = original_verify
        self.assertEqual(before_files, (notes_path.read_bytes(), comments_path.read_bytes()))
        with self.store._session() as db:
            after_db = [tuple(row) for row in db.execute(
                "SELECT comment_id,content,comment_status,is_deleted FROM comments WHERE note_id=? ORDER BY comment_id",
                (note["noteId"],),
            ).fetchall()]
        self.assertEqual(before_db, after_db)

    def test_access_status_rejects_cross_store_field_drift(self):
        note, old, kept = self._seed_pulled_note_with_comments()
        notes_path, _comments_path = self._configure_csv("access-field-gate")
        self.store._sync_pull_to_xlsx(note, [old, kept], {"folder": "", "files": []})
        headers, rows = self.store._read_csv_table(notes_path, NOTE_CSV_HEADERS)
        rows[0]["笔记标题"] = "外部改坏的标题"
        self.store._replace_csv_table(notes_path, headers, rows, "simulate-field-drift")
        with self.store._session() as db:
            before = tuple(db.execute(
                "SELECT access_status,post_status,is_deleted FROM notes WHERE note_id=?", (note["noteId"],)
            ).fetchone())
        with self.assertRaisesRegex(ValueError, "字段一致性"):
            self.store.set_note_access_status({"noteId": note["noteId"], "status": "unreachable"})
        with self.store._session() as db:
            after = tuple(db.execute(
                "SELECT access_status,post_status,is_deleted FROM notes WHERE note_id=?", (note["noteId"],)
            ).fetchone())
            active_comments = db.execute(
                "SELECT COUNT(*) FROM comments WHERE note_id=? AND is_deleted=0", (note["noteId"],)
            ).fetchone()[0]
        self.assertEqual(before, after)
        self.assertEqual(2, active_comments)

    def test_failed_pull_restores_managed_media_files_and_snapshots(self):
        note, old, kept = self._seed_pulled_note_with_comments()
        notes_path, comments_path = self._configure_csv("media-rollback")
        media_dir = Path(self.tmp.name) / "posts_materials" / "rollback"
        media_dir.mkdir(parents=True)
        (media_dir / "image-01.jpg").write_bytes(b"original-image")
        with self.store._session() as db:
            db.execute("UPDATE notes SET media_dir=? WHERE note_id=?", (str(media_dir), note["noteId"]))
        self.store._sync_pull_to_xlsx(
            note, [old, kept], {"folder": str(media_dir), "files": ["image-01.jpg"]}
        )
        self.store._refresh_material_snapshot_for_note(note["noteId"])
        before_files = {
            path.name: path.read_bytes() for path in media_dir.iterdir() if path.is_file()
        }
        before_csv = (notes_path.read_bytes(), comments_path.read_bytes())
        original_download = self.store._download_note_media
        original_verify = self.store._verify_note_store_consistency

        def destructive_download(_note):
            (media_dir / "image-01.jpg").write_bytes(b"changed-image")
            (media_dir / "image-02.jpg").write_bytes(b"new-image")
            (media_dir / "note.json").write_text("{}", encoding="utf-8")
            return {
                "status": "complete", "folder": str(media_dir),
                "files": ["image-01.jpg", "image-02.jpg", "note.json", "帖子正文.txt"],
                "fileCount": 4, "imageCount": 2, "videoCount": 0, "error": "",
            }

        self.store._download_note_media = destructive_download
        self.store._verify_note_store_consistency = lambda *_args, **_kwargs: (_ for _ in ()).throw(
            ValueError("forced media rollback")
        )
        try:
            with self.assertRaisesRegex(ValueError, "forced media rollback"):
                self.store.pull_to_excel({
                    "note": note, "comments": [old, kept], "expectedCount": 2,
                    "commentStatus": "likely_complete",
                })
        finally:
            self.store._download_note_media = original_download
            self.store._verify_note_store_consistency = original_verify
        after_files = {path.name: path.read_bytes() for path in media_dir.iterdir() if path.is_file()}
        self.assertEqual(before_files, after_files)
        self.assertEqual(before_csv, (notes_path.read_bytes(), comments_path.read_bytes()))

    def test_legacy_comment_url_and_deleted_status_schema_migration_is_lossless(self):
        notes_path, comments_path = self._configure_csv("legacy-comment-columns")
        note = {
            "noteId": "legacynote123", "title": "旧评论列迁移", "content": "正文",
            "author": "帖主", "authorUrl": "https://www.xiaohongshu.com/user/profile/post-owner",
            "url": "https://www.xiaohongshu.com/explore/legacynote123", "detailRead": True,
        }
        existing = {
            "commentId": "legacy-url-existing", "author": "评论者甲", "content": "旧列评论甲",
            "authorUrl": "https://www.xiaohongshu.com/user/profile/comment-owner-a",
        }
        self.store.confirm(note)
        self.store.upsert_comments({"noteId": note["noteId"], "comments": [existing], "status": "partial"})
        self.store._sync_pull_to_xlsx(note, [existing], {"folder": "", "files": []})
        with self.store._session() as db:
            db.execute(
                """UPDATE comments SET comment_status='已删除',is_deleted=1,
                   author_url=? WHERE comment_id=?""",
                (existing["authorUrl"], existing["commentId"]),
            )
        legacy_headers = [
            name for name in COMMENT_CSV_HEADERS if name not in {"评论用户主页url", "评论状态"}
        ]
        legacy_rows = [
            {
                "笔记ID": note["noteId"], "原笔记url": note["url"],
                "帖子用户主页url": existing["authorUrl"], "笔记评论ID": existing["commentId"],
                "用户昵称": existing["author"], "评论内容": existing["content"],
            },
            {
                "笔记ID": note["noteId"], "原笔记url": note["url"],
                "帖子用户主页url": "https://www.xiaohongshu.com/user/profile/comment-owner-b",
                "笔记评论ID": "legacy-url-new", "用户昵称": "评论者乙", "评论内容": "旧列评论乙",
            },
        ]
        self.store._replace_csv_table(comments_path, legacy_headers, legacy_rows, "legacy-schema")

        self.store._ensure_seed_workbook(notes_path)
        _headers, migrated = self.store._read_csv_table(comments_path, COMMENT_CSV_HEADERS)
        by_id = {row["笔记评论ID"]: row for row in migrated}
        self.assertEqual(existing["authorUrl"], by_id[existing["commentId"]]["评论用户主页url"])
        self.assertEqual("https://www.xiaohongshu.com/user/profile/comment-owner-b",
                         by_id["legacy-url-new"]["评论用户主页url"])
        self.assertEqual(note["authorUrl"], by_id[existing["commentId"]]["帖子用户主页url"])
        self.assertEqual("已删除", by_id[existing["commentId"]]["评论状态"])
        self.store.seed_from_xlsx(notes_path)
        with self.store._session() as db:
            saved = db.execute(
                "SELECT comment_id,author_url,comment_status,is_deleted FROM comments "
                "WHERE comment_id IN ('legacy-url-existing','legacy-url-new') ORDER BY comment_id"
            ).fetchall()
        self.assertEqual(("legacy-url-existing", existing["authorUrl"], "已删除", 1), tuple(saved[0]))
        self.assertEqual(("legacy-url-new", "https://www.xiaohongshu.com/user/profile/comment-owner-b", "存在", 0),
                         tuple(saved[1]))

    def test_empty_comment_snapshot_requires_explicit_evidence_and_deduplicated_count(self):
        note, old, kept = self._seed_pulled_note_with_comments()
        self._configure_csv("empty-evidence")
        self.store._sync_pull_to_xlsx(note, [old, kept], {"folder": "", "files": []})
        unverified = self.store.sync_comment_snapshot({
            "noteId": note["noteId"], "note": note, "comments": [],
            "expectedCount": 0, "status": "likely_complete",
        })
        self.assertFalse(unverified["canPrune"])
        self.assertEqual(0, unverified["commentsMarkedDeleted"])
        duplicate = self.store.compare_comments({
            "noteId": note["noteId"], "comments": [kept, dict(kept)],
            "expectedCount": 2, "collectionEvidence": COMPLETE_EVIDENCE,
            "status": "likely_complete",
        })
        self.assertEqual(1, duplicate["currentCount"])
        self.assertFalse(duplicate["canPrune"])
        verified = self.store.sync_comment_snapshot({
            "noteId": note["noteId"], "note": note, "comments": [],
            "expectedCount": 0, "explicitEmptyVerified": True,
            "collectionEvidence": COMPLETE_EVIDENCE, "status": "likely_complete",
        })
        self.assertTrue(verified["canPrune"])
        self.assertEqual(2, verified["commentsMarkedDeleted"])

    def test_access_status_without_exact_business_csv_row_rolls_back(self):
        notes_path, _comments_path = self._configure_csv("access-no-row")
        note = {
            "noteId": "accessnorow123", "title": "不在业务表", "content": "正文",
            "url": "https://www.xiaohongshu.com/explore/accessnorow123",
        }
        self.store.confirm(note)
        before_csv = notes_path.read_bytes()
        with self.assertRaisesRegex(ValueError, "恰有一行"):
            self.store.set_note_access_status({"noteId": note["noteId"], "status": "unreachable"})
        self.assertEqual(before_csv, notes_path.read_bytes())
        with self.store._session() as db:
            saved = db.execute(
                "SELECT access_status,post_status,is_deleted FROM notes WHERE note_id=?", (note["noteId"],)
            ).fetchone()
        self.assertEqual(("", "存在", 0), tuple(saved))

    def test_partial_media_result_rolls_back_old_binary_instead_of_committing(self):
        note, old, kept = self._seed_pulled_note_with_comments()
        notes_path, comments_path = self._configure_csv("partial-media-rollback")
        media_dir = Path(self.tmp.name) / "posts_materials" / "partial-media"
        media_dir.mkdir(parents=True)
        (media_dir / "image-01.jpg").write_bytes(b"original")
        with self.store._session() as db:
            db.execute("UPDATE notes SET media_dir=? WHERE note_id=?", (str(media_dir), note["noteId"]))
        self.store._sync_pull_to_xlsx(note, [old, kept], {"folder": str(media_dir), "files": ["image-01.jpg"]})
        self.store._refresh_material_snapshot_for_note(note["noteId"])
        before = {path.name: path.read_bytes() for path in media_dir.iterdir() if path.is_file()}
        before_csv = (notes_path.read_bytes(), comments_path.read_bytes())
        original_download = self.store._download_note_media

        def partial_download(_note):
            (media_dir / "image-01.jpg").write_bytes(b"damaged")
            (media_dir / "image-02.jpg").write_bytes(b"partial-new")
            return {"status": "partial", "folder": str(media_dir), "files": ["image-01.jpg"],
                    "fileCount": 1, "imageCount": 1, "videoCount": 0, "error": "图片2下载失败"}

        self.store._download_note_media = partial_download
        try:
            with self.assertRaisesRegex(ValueError, "素材下载未完整"):
                self.store.pull_to_excel({
                    "note": note, "comments": [old, kept], "expectedCount": 2,
                    "commentStatus": "likely_complete",
                })
        finally:
            self.store._download_note_media = original_download
        after = {path.name: path.read_bytes() for path in media_dir.iterdir() if path.is_file()}
        self.assertEqual(before, after)
        self.assertEqual(before_csv, (notes_path.read_bytes(), comments_path.read_bytes()))

    def test_persistent_checkpoint_recovers_interrupted_sync_on_next_start(self):
        notes_path, comments_path = self._configure_csv("persistent-checkpoint")
        note = {
            "noteId": "recovernote123", "title": "崩溃恢复", "content": "原正文",
            "url": "https://www.xiaohongshu.com/explore/recovernote123", "detailRead": True,
        }
        comment = {"commentId": "recover-comment", "author": "用户", "content": "原评论"}
        self.store.confirm(note)
        self.store.upsert_comments({"noteId": note["noteId"], "comments": [comment], "status": "partial"})
        media_dir = Path(self.tmp.name) / "posts_materials" / "recover"
        media_dir.mkdir(parents=True)
        (media_dir / "image-01.jpg").write_bytes(b"before-crash")
        with self.store._session() as db:
            db.execute("UPDATE notes SET media_dir=? WHERE note_id=?", (str(media_dir), note["noteId"]))
        self.store._sync_pull_to_xlsx(note, [comment], {"folder": str(media_dir), "files": ["image-01.jpg"]})
        self.store._refresh_material_snapshot_for_note(note["noteId"])
        before_csv = (notes_path.read_bytes(), comments_path.read_bytes())
        checkpoint = self.store._capture_sync_checkpoint(note["noteId"], capture_media_binaries=True)
        with self.store._session() as db:
            db.execute("UPDATE comments SET content='崩溃后的半成品' WHERE comment_id='recover-comment'")
        (media_dir / "image-01.jpg").write_bytes(b"after-crash")
        self.store._replace_csv_table(notes_path, NOTE_CSV_HEADERS, [], "interrupted")
        self.assertTrue(Path(checkpoint["checkpointDir"], "checkpoint.json").is_file())

        recovered = MonitorStore(self.store.db_path, self.store.export_dir, ai_client=FakeAI())
        recovered.configure_data_files(notes_path)
        self.assertEqual(before_csv, (notes_path.read_bytes(), comments_path.read_bytes()))
        with recovered._session() as db:
            content = db.execute(
                "SELECT content FROM comments WHERE comment_id='recover-comment'"
            ).fetchone()[0]
        self.assertEqual("原评论", content)
        self.assertEqual(b"before-crash", (media_dir / "image-01.jpg").read_bytes())
        self.assertFalse((self.store.export_dir / ".sync_checkpoints").exists())

    def test_manual_review_write_waits_for_transactional_comment_sync(self):
        self._configure_csv("concurrent-review")
        note = {
            "noteId": "concurrent123", "title": "并发串行", "content": "正文",
            "url": "https://www.xiaohongshu.com/explore/concurrent123", "detailRead": True,
        }
        comment = {"commentId": "concurrent-comment", "author": "用户", "content": "评论"}
        self.store.confirm(note)
        self.store.upsert_comments({"noteId": note["noteId"], "comments": [comment], "status": "partial"})
        folder = Path(self.tmp.name) / "posts_materials" / "concurrent"
        folder.mkdir(parents=True)
        with self.store._session() as db:
            db.execute("UPDATE notes SET media_dir=? WHERE note_id=?", (str(folder), note["noteId"]))
        self.store._sync_pull_to_xlsx(note, [comment], {"folder": str(folder), "files": []})
        self.store._refresh_material_snapshot_for_note(note["noteId"])
        original_write = self.store._write_media_snapshot
        entered = threading.Event()
        release = threading.Event()
        review_done = threading.Event()
        errors = []

        def blocked_write(*args, **kwargs):
            entered.set()
            if not release.wait(5):
                raise TimeoutError("test release timeout")
            return original_write(*args, **kwargs)

        def sync_worker():
            try:
                self.store.sync_comment_snapshot({
                    "noteId": note["noteId"], "note": note, "comments": [comment],
                    "expectedCount": 1, "collectionEvidence": COMPLETE_EVIDENCE,
                    "status": "likely_complete",
                })
            except Exception as exc:
                errors.append(exc)

        def review_worker():
            try:
                self.store.update_review({
                    "targetType": "note", "targetId": note["noteId"],
                    "reviewStatus": "resolved", "manualNegative": True,
                })
            except Exception as exc:
                errors.append(exc)
            finally:
                review_done.set()

        self.store._write_media_snapshot = blocked_write
        sync_thread = threading.Thread(target=sync_worker)
        review_thread = threading.Thread(target=review_worker)
        try:
            sync_thread.start()
            self.assertTrue(entered.wait(5))
            review_thread.start()
            self.assertFalse(review_done.wait(.2))
            release.set()
            sync_thread.join(10)
            review_thread.join(10)
        finally:
            release.set()
            self.store._write_media_snapshot = original_write
        self.assertFalse(sync_thread.is_alive())
        self.assertFalse(review_thread.is_alive())
        self.assertEqual([], errors)
        with self.store._session() as db:
            row = db.execute(
                "SELECT review_status,manual_negative FROM notes WHERE note_id=?", (note["noteId"],)
            ).fetchone()
        self.assertEqual(("resolved", 1), tuple(row))

    def test_global_repair_checkpoint_capture_failure_cleans_prior_checkpoints(self):
        notes_path, comments_path = self._configure_csv("capture-failure")
        for suffix in ("aaa", "bbb"):
            note = {
                "noteId": f"capture{suffix}123", "title": f"捕获{suffix}", "content": "正文",
                "url": f"https://www.xiaohongshu.com/explore/capture{suffix}123",
            }
            self.store.confirm(note)
            self.store._sync_pull_to_xlsx(note, [], {"folder": "", "files": []})
        before_csv = (notes_path.read_bytes(), comments_path.read_bytes())
        with self.store._session() as db:
            before_db = [tuple(row) for row in db.execute(
                "SELECT note_id,title FROM notes ORDER BY note_id"
            ).fetchall()]
        original_capture = self.store._capture_sync_checkpoint
        calls = 0

        def fail_second(note_id, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise OSError("forced checkpoint capture failure")
            return original_capture(note_id, **kwargs)

        self.store._capture_sync_checkpoint = fail_second
        try:
            with self.assertRaisesRegex(OSError, "forced checkpoint capture failure"):
                self.store.repair_csv_relationships({})
        finally:
            self.store._capture_sync_checkpoint = original_capture
        self.assertEqual(before_csv, (notes_path.read_bytes(), comments_path.read_bytes()))
        with self.store._session() as db:
            after_db = [tuple(row) for row in db.execute(
                "SELECT note_id,title FROM notes ORDER BY note_id"
            ).fetchall()]
        self.assertEqual(before_db, after_db)
        self.assertFalse((self.store.export_dir / ".sync_checkpoints").exists())

    def test_failed_global_health_repair_restores_all_checkpointed_stores(self):
        notes_path, comments_path = self._configure_csv("global-repair-rollback")
        note = {
            "noteId": "globalrepair123", "title": "全局修复回滚", "content": "正文",
            "url": "https://www.xiaohongshu.com/explore/globalrepair123", "detailRead": True,
        }
        comment = {"commentId": "global-repair-comment", "author": "用户", "content": "评论"}
        self.store.confirm(note)
        self.store.upsert_comments({"noteId": note["noteId"], "comments": [comment], "status": "partial"})
        folder = Path(self.tmp.name) / "posts_materials" / "global-repair"
        folder.mkdir(parents=True)
        (folder / "image-01.jpg").write_bytes(b"global-before")
        with self.store._session() as db:
            db.execute("UPDATE notes SET media_dir=?,comment_count_collected=99 WHERE note_id=?",
                       (str(folder), note["noteId"]))
        self.store._sync_pull_to_xlsx(note, [comment], {"folder": str(folder), "files": ["image-01.jpg"]})
        self.store._refresh_material_snapshot_for_note(note["noteId"])
        before_csv = (notes_path.read_bytes(), comments_path.read_bytes())
        before_files = {path.name: path.read_bytes() for path in folder.iterdir() if path.is_file()}
        with self.store._session() as db:
            before_db = tuple(db.execute(
                "SELECT comment_count_collected,media_dir FROM notes WHERE note_id=?", (note["noteId"],)
            ).fetchone())
        original_refresh = self.store._refresh_all_material_snapshots

        def fail_after_unrelated_database_write():
            with self.store._session() as db:
                db.execute(
                    "INSERT INTO data_repair_archive VALUES(?,?,?,?,?,?)",
                    ("must-rollback", "test", 1, "forced", "{}", "now"),
                )
            raise OSError("forced global failure")

        self.store._refresh_all_material_snapshots = fail_after_unrelated_database_write
        try:
            result = self.store.repair_data_health({})
        finally:
            self.store._refresh_all_material_snapshots = original_refresh
        self.assertFalse(result["ok"])
        self.assertTrue(result["rolledBack"])
        self.assertEqual(before_csv, (notes_path.read_bytes(), comments_path.read_bytes()))
        self.assertEqual(before_files, {path.name: path.read_bytes() for path in folder.iterdir() if path.is_file()})
        with self.store._session() as db:
            after_db = tuple(db.execute(
                "SELECT comment_count_collected,media_dir FROM notes WHERE note_id=?", (note["noteId"],)
            ).fetchone())
            unrelated_count = db.execute(
                "SELECT COUNT(*) FROM data_repair_archive WHERE repair_id='must-rollback'"
            ).fetchone()[0]
        self.assertEqual(before_db, after_db)
        self.assertEqual(0, unrelated_count)
        self.assertFalse((self.store.export_dir / ".sync_checkpoints").exists())

    def test_failed_new_pull_preserves_preexisting_unlinked_user_material_folder(self):
        notes_path, comments_path = self._configure_csv("preexisting-material")
        note = {
            "noteId": "preexisting123", "title": "预存用户目录", "content": "正文",
            "url": "https://www.xiaohongshu.com/explore/preexisting123", "detailRead": True,
        }
        folder = notes_path.parent / "posts_materials" / self.store._safe_media_folder_name(note)
        folder.mkdir(parents=True)
        (folder / "user-file.txt").write_bytes(b"must-survive")
        before_csv = (notes_path.read_bytes(), comments_path.read_bytes())
        original_download = self.store._download_note_media

        def fail_after_writing(_note):
            (folder / "image-01.jpg").write_bytes(b"transaction-file")
            (folder / "note.json").write_text("{}", encoding="utf-8")
            raise OSError("forced material failure")

        self.store._download_note_media = fail_after_writing
        try:
            with self.assertRaisesRegex(ValueError, "素材快照写入失败"):
                self.store.pull_to_excel({"note": note, "comments": [], "commentStatus": "partial"})
        finally:
            self.store._download_note_media = original_download
        self.assertTrue(folder.is_dir())
        self.assertEqual(b"must-survive", (folder / "user-file.txt").read_bytes())
        self.assertFalse((folder / "image-01.jpg").exists())
        self.assertFalse((folder / "note.json").exists())
        self.assertEqual(before_csv, (notes_path.read_bytes(), comments_path.read_bytes()))

    def test_startup_recovery_accepts_declared_internal_csv_crash_window(self):
        notes_path, _comments_path = self._configure_csv("checkpoint-internal-window")
        note = {
            "noteId": "internalwindow123", "title": "内部替换窗口", "content": "正文",
            "url": "https://www.xiaohongshu.com/explore/internalwindow123",
        }
        self.store.confirm(note)
        self.store._sync_pull_to_xlsx(note, [], {"folder": "", "files": []})
        before = notes_path.read_bytes()
        checkpoint = self.store._capture_sync_checkpoint(note["noteId"])
        headers, rows = self.store._read_csv_table(notes_path, NOTE_CSV_HEADERS)
        rows[0]["笔记标题"] = "尚未提交完成的内部标题"
        temporary = self.store._write_csv_temporary(notes_path, headers, rows, "crash-window")
        self.store._prepare_active_checkpoint_csv_replacement({notes_path: temporary})
        replace_with_retry(temporary, notes_path)
        self.assertNotEqual(before, notes_path.read_bytes())

        recovered = MonitorStore(self.store.db_path, self.store.export_dir, ai_client=FakeAI())
        recovered.configure_data_files(notes_path)
        self.assertEqual(before, notes_path.read_bytes())
        self.assertFalse(Path(checkpoint["checkpointDir"]).exists())
        self.store._active_csv_checkpoint = None

    def test_startup_recovery_stops_on_external_csv_conflict_and_preserves_copy(self):
        notes_path, _comments_path = self._configure_csv("checkpoint-conflict")
        note = {
            "noteId": "conflictnote123", "title": "恢复冲突", "content": "正文",
            "url": "https://www.xiaohongshu.com/explore/conflictnote123",
        }
        self.store.confirm(note)
        self.store._sync_pull_to_xlsx(note, [], {"folder": "", "files": []})
        checkpoint = self.store._capture_sync_checkpoint(note["noteId"])
        external_bytes = notes_path.read_bytes() + b"\r\n# external edit"
        notes_path.write_bytes(external_bytes)

        recovered = MonitorStore(self.store.db_path, self.store.export_dir, ai_client=FakeAI())
        with self.assertRaisesRegex(RuntimeError, "外部 CSV 修改"):
            recovered.configure_data_files(notes_path)
        self.assertEqual(external_bytes, notes_path.read_bytes())
        conflicts = list(Path(checkpoint["checkpointDir"]).glob("external-conflict-*.csv"))
        self.assertEqual(1, len(conflicts))
        self.assertEqual(external_bytes, conflicts[0].read_bytes())
        self.store._discard_sync_checkpoint(checkpoint)

    def test_verifier_rejects_duplicate_csv_rows_and_cross_note_comment_move(self):
        notes_path, comments_path = self._configure_csv("identity-guards")
        first = {
            "noteId": "guardnote123", "title": "帖子甲", "content": "甲",
            "url": "https://www.xiaohongshu.com/explore/guardnote123",
        }
        second = {
            "noteId": "guardnote456", "title": "帖子乙", "content": "乙",
            "url": "https://www.xiaohongshu.com/explore/guardnote456",
        }
        comment = {"commentId": "guard-comment-1", "author": "用户", "content": "不可串帖"}
        for note in (first, second):
            self.store.confirm(note)
            self.store._sync_pull_to_xlsx(note, [comment] if note is first else [], {"folder": "", "files": []})
        self.store.upsert_comments({"noteId": first["noteId"], "comments": [comment], "status": "partial"})
        headers, rows = self.store._read_csv_table(comments_path, COMMENT_CSV_HEADERS)
        rows.append(dict(rows[0]))
        self.store._replace_csv_table(comments_path, headers, rows, "duplicate-id")
        with self.assertRaisesRegex(ValueError, "重复评论 ID"):
            self.store._verify_note_store_consistency(first["noteId"])
        self.store._replace_csv_table(comments_path, headers, rows[:1], "remove-duplicate")
        with self.assertRaisesRegex(ValueError, "属于其他帖子"):
            self.store._sync_pull_to_xlsx(second, [comment], {"folder": "", "files": []})
        self.assertEqual(2, len(self.store._read_csv_table(notes_path, NOTE_CSV_HEADERS)[1]))
        self.assertEqual(first["noteId"], self.store._read_csv_table(comments_path, COMMENT_CSV_HEADERS)[1][0]["笔记ID"])

    def test_comment_sync_uses_one_canonical_long_title_across_stores(self):
        notes_path, _comments_path = self._configure_csv("canonical-title")
        note = {
            "noteId": "longtitle123", "title": "很长标题" * 30, "content": "正文",
            "url": "https://www.xiaohongshu.com/explore/longtitle123", "detailRead": True,
        }
        comment = {"commentId": "long-title-comment", "author": "用户", "content": "评论"}
        self.store.confirm(note)
        self.store.upsert_comments({"noteId": note["noteId"], "comments": [comment], "status": "partial"})
        self.store._sync_pull_to_xlsx(note, [comment], {"folder": "", "files": []})
        result = self.store.sync_comment_snapshot({
            "noteId": note["noteId"], "note": note, "comments": [comment],
            "expectedCount": 1, "collectionEvidence": COMPLETE_EVIDENCE,
            "status": "likely_complete",
        })
        self.assertTrue(result["consistencyVerified"])
        csv_title = self.store._read_csv_table(notes_path, NOTE_CSV_HEADERS)[1][0]["笔记标题"]
        with self.store._session() as db:
            db_title = db.execute("SELECT title FROM notes WHERE note_id=?", (note["noteId"],)).fetchone()[0]
        self.assertEqual(csv_title, db_title)
        self.assertLessEqual(len(db_title), 80)

    def test_empty_csv_source_fields_preserve_tags_and_db_analysis_repairs_csv(self):
        notes_path, comments_path = self._configure_csv("source-preserve")
        note = {
            "noteId": "preservetags123", "title": "保留来源字段", "content": "正文",
            "tags": ["#保留话题"], "url": "https://www.xiaohongshu.com/explore/preservetags123",
            "detailRead": True,
        }
        comment = {"commentId": "preserve-comment", "author": "用户", "content": "评论"}
        self.store.confirm(note)
        self.store.upsert_comments({"noteId": note["noteId"], "comments": [comment], "status": "partial"})
        self.store._sync_pull_to_xlsx(note, [comment], {"folder": "", "files": []})
        headers, rows = self.store._read_csv_table(notes_path, NOTE_CSV_HEADERS)
        rows[0]["笔记话题"] = ""
        rows[0]["AI情绪判断"] = ""
        rows[0]["帖子好坏"] = ""
        self.store._replace_csv_table(notes_path, headers, rows, "blank-source")
        with self.store._session() as db:
            db.execute("UPDATE notes SET post_sentiment='negative' WHERE note_id=?", (note["noteId"],))
        self.store._seed_from_csv(notes_path)
        with self.store._session() as db:
            stored_tags = db.execute("SELECT tags FROM notes WHERE note_id=?", (note["noteId"],)).fetchone()[0]
        self.assertEqual("#保留话题", stored_tags)
        result = self.store.sync_comment_snapshot({
            "noteId": note["noteId"], "note": {"noteId": note["noteId"], "detailRead": True},
            "comments": [comment], "expectedCount": 1, "status": "likely_complete",
        })
        self.assertTrue(result["consistencyVerified"])
        saved_note = self.store._read_csv_table(notes_path, NOTE_CSV_HEADERS)[1][0]
        self.assertEqual("#保留话题", saved_note["笔记话题"])
        self.assertEqual(("差评", "差评"), (saved_note["AI情绪判断"], saved_note["帖子好坏"]))
        self.assertEqual(1, len(self.store._read_csv_table(comments_path, COMMENT_CSV_HEADERS)[1]))

    def test_data_health_reports_comment_aliases_without_auto_merging(self):
        self._configure_csv("identity-audit")
        note = {
            "noteId": "identitynote123", "title": "评论身份体检", "content": "正文",
            "author": "作者", "url": "https://www.xiaohongshu.com/explore/identitynote123",
        }
        self.store.confirm(note)
        comments = [
            {"commentId": "raw-id-1", "author": "用户", "content": "同一条", "publishedAt": "09-01"},
            {"commentId": "comment-raw-id-1", "author": "用户", "content": "同一条", "publishedAt": "09-01"},
            {"commentId": "legacy-separate-id", "author": "用户", "content": "同一条", "publishedAt": "09-01"},
        ]
        self.store.upsert_comments({"noteId": note["noteId"], "comments": comments, "status": "partial"})
        self.store._sync_pull_to_xlsx(note, comments, {"folder": "", "files": []})

        health = self.store.data_health()
        self.assertEqual(1, health["summary"]["commentIdAliasCandidates"])
        self.assertEqual(1, health["summary"]["logicalDuplicateCommentCandidates"])
        issue_ids = {item["id"] for item in health["issues"]}
        self.assertIn("comment_id_alias_candidates", issue_ids)
        self.assertIn("comment_logical_duplicate_candidates", issue_ids)
        self.assertTrue(health["summary"]["relationshipsConsistent"])
        with self.store._session() as db:
            self.assertEqual(3, db.execute(
                "SELECT COUNT(*) FROM comments WHERE note_id=?", (note["noteId"],)
            ).fetchone()[0])

    def test_tag_and_chinese_comment_level_normalization(self):
        self.assertEqual("#Origani #护肤", tag_text("# O r i g a n 护 肤 #Origani #护肤"))
        self.assertEqual("#现代美容仪", tag_text("# 现 代 美 容 仪 作 者"))
        self.assertEqual(2, comment_level_value("二级评论"))
        notes_path, comments_path = self._configure_csv("chinese-level")
        note_headers, note_rows = self.store._read_csv_table(notes_path, NOTE_CSV_HEADERS)
        note_rows.append({**{name: "" for name in note_headers},
                          "笔记ID": "levelnote123", "笔记url": "https://www.xiaohongshu.com/explore/levelnote123",
                          "笔记标题": "层级测试", "笔记内容": "正文", "帖子状态": "存在"})
        self.store._replace_csv_table(notes_path, note_headers, note_rows, "level-note")
        comment_headers, comment_rows = self.store._read_csv_table(comments_path, COMMENT_CSV_HEADERS)
        comment_rows.append({**{name: "" for name in comment_headers},
                             "笔记ID": "levelnote123", "原笔记url": "https://www.xiaohongshu.com/explore/levelnote123",
                             "笔记评论ID": "level-comment-1", "用户昵称": "用户", "评论内容": "回复",
                             "评论层级": "二级评论", "父评论ID": "parent-1", "评论状态": "存在"})
        self.store._replace_csv_table(comments_path, comment_headers, comment_rows, "level-comment")
        self.store.seed_from_xlsx(notes_path)
        with self.store._session() as db:
            stored_level = db.execute(
                "SELECT comment_level,payload_json FROM comments WHERE comment_id='level-comment-1'"
            ).fetchone()
        self.assertEqual(2, stored_level[0])
        self.assertEqual("子评论", json.loads(stored_level[1])["commentType"])

        self.store.upsert_comments({
            "noteId": "levelnote123", "status": "partial", "comments": [{
                "commentId": "level-comment-2", "parentCommentId": "parent-2",
                "commentLevel": 1, "author": "用户2", "content": "新回复",
            }]
        })
        with self.store._session() as db:
            inserted = db.execute(
                "SELECT comment_level,payload_json FROM comments WHERE comment_id='level-comment-2'"
            ).fetchone()
        self.assertEqual(2, inserted[0])
        self.assertEqual("子评论", json.loads(inserted[1])["commentType"])

    def test_partial_sync_never_marks_unloaded_comments_deleted(self):
        note, old, kept = self._seed_pulled_note_with_comments()
        _notes_path, comments_path = self._configure_csv("partial-presence")
        self.store._sync_pull_to_xlsx(note, [old, kept], {"folder": "", "files": []})
        result = self.store.sync_comment_snapshot({
            "noteId": note["noteId"], "note": note, "comments": [kept],
            "expectedCount": 2, "status": "partial"
        })
        self.assertFalse(result["canPrune"])
        csv_rows = {row["笔记评论ID"]: row for row in self._csv_rows(comments_path, COMMENT_CSV_HEADERS)}
        self.assertEqual("存在", csv_rows["comment-old"]["评论状态"])
        with self.store._session() as db:
            deleted = db.execute(
                "SELECT is_deleted FROM comments WHERE comment_id='comment-old'"
            ).fetchone()[0]
        self.assertEqual(0, deleted)

    def test_csv_semantic_and_presence_fields_import_to_sqlite(self):
        notes_path, comments_path = self._configure_csv("semantic-import")
        note_headers, note_rows = self.store._read_csv_table(notes_path, NOTE_CSV_HEADERS)
        note_rows.append({**{name: "" for name in note_headers},
                          "笔记ID": "semanticnote123", "笔记url": "https://www.xiaohongshu.com/explore/semanticnote123",
                          "笔记标题": "语义字段帖子", "语义分析次数": "3", "分析结论是否差评": "是",
                          "差评类型": "Price discrepancy", "差评子类型": "Felt overcharged"})
        self.store._replace_csv_table(notes_path, note_headers, note_rows, "semantic-note")
        comment_headers, comment_rows = self.store._read_csv_table(comments_path, COMMENT_CSV_HEADERS)
        comment_rows.append({**{name: "" for name in comment_headers},
                             "笔记ID": "semanticnote123", "原笔记url": "https://www.xiaohongshu.com/explore/semanticnote123",
                             "笔记评论ID": "semantic-comment-1", "用户昵称": "用户", "评论内容": "历史评论",
                             "评论层级": "1级评论", "语义分析次数": "2", "分析结论是否差评": "是",
                             "差评类型": "Sales pitch", "差评子类型": "Aggressive pull-in", "评论状态": "已删除"})
        self.store._replace_csv_table(comments_path, comment_headers, comment_rows, "semantic-comment")

        self.store.seed_from_xlsx(notes_path)

        with self.store._session() as db:
            note_row = db.execute("""SELECT semantic_analysis_count,analysis_is_negative,negative_type,negative_subtype
                                      FROM notes WHERE note_id='semanticnote123'""").fetchone()
            comment_row = db.execute("""SELECT semantic_analysis_count,analysis_is_negative,negative_type,
                                         negative_subtype,comment_status,is_deleted
                                         FROM comments WHERE comment_id='semantic-comment-1'""").fetchone()
        self.assertEqual((3, "是", "Price discrepancy", "Felt overcharged"), tuple(note_row))
        self.assertEqual((2, "是", "Sales pitch", "Aggressive pull-in", "已删除", 1), tuple(comment_row))

    def test_note_status_repairs_legacy_csv_media_path(self):
        note_id = "legacy123456"
        notes_path, _comments_path = self._configure_csv("legacy")
        actual_folder = notes_path.parent / "posts_materials" / "旧标题"
        actual_folder.mkdir(parents=True)
        (actual_folder / "旧标题-图1.jpg").write_bytes(b"image")
        headers, rows = self.store._read_csv_table(notes_path, NOTE_CSV_HEADERS)
        row = {name: "" for name in headers}
        row.update({"笔记ID": note_id, "笔记标题": "旧标题",
                    "笔记url": f"https://www.xiaohongshu.com/explore/{note_id}",
                    "对应帖子文件夹地址": "C:/removed-root/posts_materials/旧标题",
                    "文件夹内清单": "旧标题-图1.jpg"})
        rows.append(row)
        self.store._replace_csv_table(notes_path, headers, rows, "test")
        self.store.seed_from_xlsx(notes_path)
        result = self.store.note_status(note_id)
        self.assertTrue(result["inExcel"])
        self.assertEqual(str(actual_folder.resolve()), result["mediaDir"])
        self.assertEqual(["旧标题-图1.jpg"], result["mediaFiles"])
        self.assertGreater(result["excelRow"], 1)

    def test_csv_artifact_index_is_reused_until_file_changes(self):
        from unittest.mock import patch
        note_id = "cache123456"
        notes_path, _comments_path = self._configure_csv("cache")
        headers, rows = self.store._read_csv_table(notes_path, NOTE_CSV_HEADERS)
        row = {name: "" for name in headers}
        row.update({"笔记ID": note_id, "对应帖子文件夹地址": "C:/materials/cache123456",
                    "文件夹内清单": "image-01.jpg"})
        rows.append(row)
        self.store._replace_csv_table(notes_path, headers, rows, "test")

        with patch.object(self.store, "_read_csv_table", wraps=self.store._read_csv_table) as mocked_read:
            first = self.store._excel_note_artifacts(note_id)
            second = self.store._excel_note_artifacts(note_id)
            self.assertEqual(first, second)
            self.assertEqual(1, mocked_read.call_count)

            with notes_path.open("r", encoding="utf-8-sig", newline="") as stream:
                raw_rows = list(csv.DictReader(stream))
            raw_rows[0]["文件夹内清单"] = "image-01.jpg\nimage-02.jpg"
            with notes_path.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=headers, lineterminator="\r\n")
                writer.writeheader(); writer.writerows(raw_rows)
            refreshed = self.store._excel_note_artifacts(note_id)
            self.assertEqual(["image-01.jpg", "image-02.jpg"], refreshed[1])
            self.assertEqual(2, mocked_read.call_count)

    def test_extension_broadcasts_delete_state_and_reconciles_unchanged_syncs(self):
        extension = Path(__file__).resolve().parent.parent / "extension"
        worker = (extension / "service-worker.js").read_text(encoding="utf-8")
        content = (extension / "content.js").read_text(encoding="utf-8")
        panel = (extension / "sidepanel.js").read_text(encoding="utf-8")
        self.assertIn("deletePulledNoteAndBroadcast", worker)
        self.assertIn('type: "localNoteStateChanged"', worker)
        self.assertIn("帖子与评论无变化，正在校准存续状态、CSV、SQLite 与素材快照", worker)
        self.assertIn("enforceExactNoteIdentity", worker)
        self.assertIn("identityConflictBlocked", content)
        self.assertIn("processPanelMutationIsInternal", content)
        self.assertIn("outsideThreshold", content)
        self.assertIn("_appliedGeometry", content)
        self.assertIn("batchFailureRenderSignature", panel)
        self.assertIn("ignoredNotes", worker)
        self.assertIn("ignoredReconciled", worker)
        self.assertIn("restoreNote(message.note", worker)
        self.assertIn('message.type === "localNoteStateChanged"', content)
        self.assertIn('message.type === "localNoteStateChanged"', panel)

    def test_powershell_output_decoder_accepts_utf8_and_utf16(self):
        self.assertEqual("WPS|True\r\n", _decode_powershell_output(b"WPS|True\r\n"))
        self.assertEqual("WPS|True\r\n", _decode_powershell_output("WPS|True\r\n".encode("utf-16le")))

    def test_atomic_replace_retries_transient_wps_lock(self):
        from unittest.mock import patch
        source = Path(self.tmp.name) / "source.xlsx"
        target = Path(self.tmp.name) / "target.xlsx"
        source.write_bytes(b"new")
        target.write_bytes(b"old")
        with patch("server.os.replace", side_effect=[PermissionError("busy"), PermissionError("busy"), None]) as replace, \
             patch("server._reopen_saved_office_workbook_read_only", return_value=False), \
             patch("server.time.sleep") as sleep:
            attempts = replace_with_retry(source, target, attempts=4, initial_delay=0.01)
        self.assertEqual(3, attempts)
        self.assertEqual(3, replace.call_count)
        self.assertEqual(2, sleep.call_count)

    def test_atomic_replace_converts_saved_wps_lock_then_retries(self):
        from unittest.mock import patch
        source = Path(self.tmp.name) / "source-wps.xlsx"
        target = Path(self.tmp.name) / "target-wps.xlsx"
        source.write_bytes(b"new")
        target.write_bytes(b"old")
        with patch("server.os.replace", side_effect=[PermissionError("wps busy"), None]) as replace, \
             patch("server._reopen_saved_office_workbook_read_only", return_value=True) as recover, \
             patch("server.time.sleep") as sleep:
            attempts = replace_with_retry(source, target, attempts=3)
        self.assertEqual(2, attempts)
        recover.assert_called_once_with(target)
        self.assertEqual(2, replace.call_count)
        sleep.assert_not_called()

    def test_open_excel_prefers_wps_and_locates_note(self):
        import base64
        import subprocess
        from unittest.mock import patch

        note_id = "openexcel123456"
        workbook_path = self.store.export_dir.parent / "wps-master.xlsx"
        self.store.seed_xlsx_path = workbook_path
        self.store._ensure_seed_workbook(workbook_path)
        pulled = self.store.pull_to_excel({
            "note": {"noteId": note_id, "title": "WPS 定位测试", "content": "正文", "detailRead": True},
            "comments": [],
        })
        completed = subprocess.CompletedProcess([], 0, stdout="WPS|True\r\n".encode("utf-16le"), stderr=b"")
        with patch("server.subprocess.run", return_value=completed) as run:
            result = self.store.open_local_artifact({
                "kind": "excel", "noteId": note_id, "excelRow": pulled["excelRow"], "fieldName": "笔记标题"
            })
        encoded_script = run.call_args.args[0][-1]
        script = base64.b64decode(encoded_script).decode("utf-16le")
        self.assertLess(script.index("ket.Application"), script.index("Excel.Application"))
        self.assertIn("$app.Workbooks.Open($path, 0, $true)", script)
        self.assertIn("$app.Goto($cell, $true)", script)
        self.assertEqual("WPS", result["application"])
        self.assertEqual("csv", result["kind"])
        self.assertTrue(result["target"].endswith(".csv"))
        self.assertIn("$book.Worksheets.Item(1)", script)
        self.assertTrue(result["readOnly"])
        self.assertEqual(pulled["excelRow"], result["row"])

    def test_open_local_artifact_uses_hydrated_media_folder(self):
        from unittest.mock import patch
        note_id = "openmedia123456"
        workbook_path = self.store.export_dir.parent / "open-master.xlsx"
        self.store.seed_xlsx_path = workbook_path
        folder = workbook_path.parent / "posts_materials" / f"素材__{note_id}"
        folder.mkdir(parents=True)
        (folder / "image-01.jpg").write_bytes(b"image")
        timestamp = "2026-08-24T15:00:00+08:00"
        with self.store._session() as db:
            db.execute("""
                INSERT INTO notes
                (note_id,url,title,first_seen_at,last_seen_at,status,is_relevant,source,
                 title_key,content_key,title_content_key,pull_status,media_dir)
                VALUES (?,?,?,?,?,'known',1,'existing_xlsx',?,'','','synced',?)
            """, (note_id, f"https://www.xiaohongshu.com/explore/{note_id}", "素材帖子",
                  timestamp, timestamp, "素材帖子", str(folder)))
        with patch("server.os.startfile") as startfile:
            result = self.store.open_local_artifact({"kind": "folder", "noteId": note_id})
        self.assertTrue(result["ok"])
        self.assertEqual(str(folder.resolve()), result["target"])
        startfile.assert_called_once_with(str(folder.resolve()))

    def test_unreachable_status_round_trip_and_bulk_mark_keeps_all_history(self):
        note_id = "unreachable123456"
        notes_path, comments_path = self._configure_csv("unreachable")
        media_dir = Path(self.tmp.name) / "posts_materials" / note_id
        media_dir.mkdir(parents=True)
        (media_dir / "image-01.jpg").write_bytes(b"image")
        note = {
            "noteId": note_id,
            "url": f"https://www.xiaohongshu.com/explore/{note_id}",
            "title": "无法打开的帖子",
            "content": "用于访问状态测试",
            "author": "测试用户",
            "detailRead": True,
        }
        comment = {
            "commentId": "comment-unreachable-1",
            "author": "评论用户",
            "content": "需要保留的历史评论",
            "publishedAt": "08-25",
        }
        self.store.confirm(note)
        self.store.upsert_comments({
            "noteId": note_id, "comments": [comment], "expectedCount": 1, "status": "likely_complete",
        })
        with self.store._session() as db:
            db.execute("""UPDATE notes SET media_dir=?,media_status='complete',media_file_count=1
                       WHERE note_id=?""", (str(media_dir), note_id))
        self.store._sync_pull_to_xlsx(note, [comment], {"folder": str(media_dir), "files": ["image-01.jpg"]})
        self.store._write_media_snapshot(
            {"folder": str(media_dir), "files": ["image-01.jpg"]}, note,
            self.store._comments_as_api(note_id),
        )
        note_headers, note_rows = self.store._read_csv_table(notes_path, NOTE_CSV_HEADERS)
        before_row = next(row for row in note_rows if row["笔记ID"] == note_id)
        before_row.update({"语义分析次数": "3", "分析结论是否差评": "是",
                           "差评类型": "Price discrepancy", "差评子类型": "Felt overcharged"})
        self.store._replace_csv_table(notes_path, note_headers, note_rows, "semantic-post")
        with self.store._session() as db:
            db.execute("""UPDATE notes SET semantic_analysis_count=3,analysis_is_negative='是',
                       negative_type='Price discrepancy',negative_subtype='Felt overcharged' WHERE note_id=?""", (note_id,))

        marked = self.store.set_note_access_status({
            "noteId": note_id, "status": "unreachable", "error": "详情页无法加载",
        })
        self.assertEqual(("unreachable", "已删除", True),
                         (marked["accessStatus"], marked["postStatus"], marked["isDeleted"]))
        note_rows = self._csv_rows(notes_path, NOTE_CSV_HEADERS)
        matching_row = next(row for row in note_rows if row["笔记ID"] == note_id)
        self.assertEqual(("打不开", "已删除"), (matching_row["访问状态"], matching_row["帖子状态"]))
        self.assertEqual(("3", "是", "Price discrepancy", "Felt overcharged"),
                         tuple(matching_row[name] for name in
                               ("语义分析次数", "分析结论是否差评", "差评类型", "差评子类型")))
        self.assertEqual(note_id, self.store.list_unreachable_notes()[0]["note_id"])
        material_note = json.loads((media_dir / "note.json").read_text(encoding="utf-8"))
        self.assertTrue(material_note["isDeleted"])
        self.assertEqual("已删除", material_note["postStatus"])
        comment_csv = next(row for row in self._csv_rows(comments_path, COMMENT_CSV_HEADERS)
                           if row["笔记评论ID"] == comment["commentId"])
        self.assertEqual("已删除", comment_csv["评论状态"])
        with self.store._session() as db:
            self.assertEqual(("已删除", 1), tuple(db.execute(
                "SELECT comment_status,is_deleted FROM comments WHERE comment_id=?", (comment["commentId"],)
            ).fetchone()))
        material_comments = json.loads((media_dir / "comments.json").read_text(encoding="utf-8"))
        self.assertTrue(material_comments[0]["isDeleted"])

        cleared = self.store.set_note_access_status({"noteId": note_id, "status": "ok"})
        self.assertEqual(("ok", "存在", False),
                         (cleared["accessStatus"], cleared["postStatus"], cleared["isDeleted"]))
        self.assertEqual([], self.store.list_unreachable_notes())
        headers, note_rows = self.store._read_csv_table(notes_path, NOTE_CSV_HEADERS)
        row = next(row for row in note_rows if row["笔记ID"] == note_id)
        row["访问状态"] = "打不开"
        row["帖子状态"] = "存在"
        self.store._replace_csv_table(notes_path, headers, note_rows, "manual")
        self.store.seed_from_xlsx(notes_path)
        self.assertEqual([], self.store.list_unreachable_notes())
        with self.store._session() as db:
            self.assertEqual(("check_failed", "存在", 0), tuple(db.execute(
                "SELECT access_status,post_status,is_deleted FROM notes WHERE note_id=?", (note_id,)
            ).fetchone()))
        migrated = self.store.reconcile_legacy_access_statuses()
        self.assertEqual(1, migrated["updated"])

        self.store.set_note_access_status({
            "noteId": note_id, "status": "unreachable", "error": "两条独立证据确认内容已删除",
        })
        with self.store._session() as db:
            db.execute(
                "INSERT INTO note_summaries(note_id,model,result_json,comment_count,created_at,updated_at) VALUES(?,?,?,?,?,?)",
                (note_id, "test", "{}", 1, "2026-08-26T16:00:00+08:00", "2026-08-26T16:00:00+08:00"),
            )
            db.execute(
                "INSERT INTO reply_generation_history(note_id,comment_id,generated_reply,created_at) VALUES(?,?,?,?)",
                (note_id, comment["commentId"], "测试回复", "2026-08-26T16:00:00+08:00"),
            )
            db.execute(
                "INSERT INTO ai_analysis_records(target_type,target_id,status,created_at) VALUES('relevance',?,'completed',?)",
                (note_id, "2026-08-26T16:00:00+08:00"),
            )
            db.execute(
                "INSERT INTO watchlist(note_id,created_at,updated_at) VALUES(?,?,?)",
                (note_id, "2026-08-26T16:00:00+08:00", "2026-08-26T16:00:00+08:00"),
            )

        result = self.store.delete_unreachable_notes({})
        self.assertTrue(result["ok"])
        self.assertTrue(result["nonDestructive"])
        self.assertEqual((1, 0), (result["markedDeletedCount"], result["deletedDatabaseComments"]))
        with self.store._session() as db:
            self.assertEqual(("已删除", 1), tuple(db.execute(
                "SELECT post_status,is_deleted FROM notes WHERE note_id=?", (note_id,)
            ).fetchone()))
            self.assertEqual(1, db.execute(
                "SELECT COUNT(*) FROM comments WHERE note_id=?", (note_id,)
            ).fetchone()[0])
            for table in ("note_summaries", "reply_generation_history", "watchlist"):
                self.assertEqual(1, db.execute(f"SELECT COUNT(*) FROM {table} WHERE note_id=?", (note_id,)).fetchone()[0])
            self.assertEqual(1, db.execute(
                "SELECT COUNT(*) FROM ai_analysis_records WHERE target_id=?", (note_id,)
            ).fetchone()[0])
        note_ids = {row["笔记ID"] for row in self._csv_rows(notes_path, NOTE_CSV_HEADERS)}
        comment_ids = {row["笔记评论ID"] for row in self._csv_rows(comments_path, COMMENT_CSV_HEADERS)}
        self.assertIn(note_id, note_ids)
        self.assertIn(comment["commentId"], comment_ids)
        self.assertTrue(media_dir.is_dir())
        health = self.store.data_health()
        self.assertTrue(health["summary"]["relationshipsConsistent"])
        self.assertEqual(1, health["summary"]["deletedPosts"])

    def test_relationship_repair_rebuilds_fragments_and_marks_orphans(self):
        known_id = "knownnote123"
        orphan_id = "orphannote123"
        notes = [
            {"笔记ID": known_id, "笔记url": f"https://www.xiaohongshu.com/explore/{known_id}", "笔记标题": "有效帖子"},
            {"笔记ID": "", "笔记url": f"https://www.xiaohongshu.com/explore/{known_id}", "笔记标题": ""},
        ]
        logical = [
            {"原笔记url": f"https://www.xiaohongshu.com/explore/{known_id}", "用户昵称": "甲",
             "评论内容": "完整保留一", "评论时间": "08-27", "评论层级": "1级评论", "点赞量": "1",
             "是否帖主评论": "否", "帖子用户主页url": "https://example.com/a", "文件夹内清单": "a.jpg"},
            {"原笔记url": f"https://www.xiaohongshu.com/explore/{orphan_id}", "用户昵称": "乙",
             "评论内容": "完整保留二", "评论时间": "08-27", "评论层级": "2级评论", "点赞量": "0",
             "是否帖主评论": "否", "帖子用户主页url": "https://example.com/b", "文件夹内清单": "b.jpg"},
        ]
        fragmented = []
        for source in logical:
            for name in ("原笔记url", "帖子用户主页url", "用户昵称", "评论内容", "评论时间",
                         "是否帖主评论", "点赞量", "评论层级", "对应帖子文件夹地址", "文件夹内清单", "AI情绪判断"):
                fragmented.append({name: source.get(name) or ("中立" if name == "AI情绪判断" else "素材" if name == "对应帖子文件夹地址" else "")})
        result = repair_relationship_rows(notes, fragmented, NOTE_CSV_HEADERS, COMMENT_CSV_HEADERS)
        self.assertEqual(1, len(result["noteRows"]))
        self.assertEqual(2, len(result["commentRows"]))
        self.assertEqual(2, result["summary"]["fragmentedCommentsRebuilt"])
        self.assertEqual(2, result["summary"]["generatedCommentIds"])
        self.assertEqual(1, result["summary"]["mappingReviewComments"])
        self.assertEqual({"完整保留一", "完整保留二"}, {row["评论内容"] for row in result["commentRows"]})
        self.assertEqual(2, len({row["笔记评论ID"] for row in result["commentRows"]}))
        orphan = next(row for row in result["commentRows"] if row["笔记ID"] == orphan_id)
        self.assertEqual("待复核", orphan["映射状态"])

    def test_relationship_repair_unions_csv_sqlite_and_material_comments(self):
        notes_path, comments_path = self._configure_csv("repair-union")
        note = {"noteId": "repairunion123", "title": "关系修复", "content": "正文",
                "url": "https://www.xiaohongshu.com/explore/repairunion123", "detailRead": True}
        self.store.confirm(note)
        folder = Path(self.tmp.name) / "posts_materials" / "repairunion123"
        folder.mkdir(parents=True)
        with self.store._session() as db:
            db.execute("UPDATE notes SET media_dir=?,source='existing_xlsx',pull_status='synced' WHERE note_id=?",
                       (str(folder), note["noteId"]))
        self.store._sync_pull_to_xlsx(note, [], {"folder": str(folder), "files": []})
        self.store.upsert_comments({
            "noteId": note["noteId"], "status": "partial",
            "comments": [{"commentId": "db-only-comment", "author": "数据库", "content": "数据库独有"}],
        })
        (folder / "comments.json").write_text(json.dumps([
            {"commentId": "material-only-comment", "noteId": note["noteId"],
             "author": "素材", "content": "素材独有", "commentLevel": 1}
        ], ensure_ascii=False), encoding="utf-8")
        legacy = {"原笔记url": note["url"], "用户昵称": "历史", "评论内容": "CSV 独有",
                  "评论时间": "08-27", "评论层级": "1级评论"}
        self.store._replace_csv_table(comments_path, COMMENT_CSV_HEADERS, [legacy], "legacy")

        result = self.store.repair_csv_relationships({})

        self.assertTrue(result["ok"])
        self.assertEqual(3, result["summary"]["commentsAfterUnion"])
        rows = self._csv_rows(comments_path, COMMENT_CSV_HEADERS)
        self.assertEqual(3, len(rows))
        self.assertTrue(all(row["笔记ID"] == note["noteId"] for row in rows))
        self.assertTrue(all(row["笔记评论ID"] for row in rows))
        with self.store._session() as db:
            self.assertEqual(3, db.execute("SELECT COUNT(*) FROM comments WHERE note_id=?", (note["noteId"],)).fetchone()[0])
        material = json.loads((folder / "comments.json").read_text(encoding="utf-8"))
        self.assertEqual({row["笔记评论ID"] for row in rows}, {item["commentId"] for item in material})
        self.assertTrue(result["health"]["summary"]["relationshipsConsistent"])

    def test_relationship_repair_splits_shared_material_directories(self):
        self._configure_csv("shared-material")
        shared = Path(self.tmp.name) / "posts_materials" / "共享目录"
        shared.mkdir(parents=True)
        (shared / "image-01.jpg").write_bytes(b"shared-image")
        notes = [
            {"noteId": "sharednote123", "title": "共享甲", "content": "甲正文",
             "url": "https://www.xiaohongshu.com/explore/sharednote123", "detailRead": True},
            {"noteId": "sharednote456", "title": "共享乙", "content": "乙正文",
             "url": "https://www.xiaohongshu.com/explore/sharednote456", "detailRead": True},
        ]
        for index, note in enumerate(notes, 1):
            comment = {"commentId": f"shared-comment-{index}", "author": f"用户{index}",
                       "content": f"评论{index}", "commentLevel": 1}
            self.store.confirm(note)
            with self.store._session() as db:
                db.execute("UPDATE notes SET media_dir=?,source='existing_xlsx',pull_status='synced' WHERE note_id=?",
                           (str(shared), note["noteId"]))
            self.store.upsert_comments({"noteId": note["noteId"], "comments": [comment], "status": "partial"})
            self.store._sync_pull_to_xlsx(note, [comment], {"folder": str(shared), "files": ["image-01.jpg"]})

        result = self.store.repair_csv_relationships({})

        self.assertEqual(2, result["summary"]["sharedMediaRecordsSplit"])
        with self.store._session() as db:
            paths = [Path(row[0]) for row in db.execute(
                "SELECT media_dir FROM notes WHERE note_id IN ('sharednote123','sharednote456') ORDER BY note_id"
            ).fetchall()]
        self.assertEqual(2, len({str(path) for path in paths}))
        self.assertFalse(shared.exists())
        for index, path in enumerate(paths, 1):
            self.assertTrue((path / "image-01.jpg").is_file())
            comments = json.loads((path / "comments.json").read_text(encoding="utf-8"))
            self.assertEqual(1, len(comments))
        self.assertTrue(result["health"]["summary"]["relationshipsConsistent"])

    def test_no_change_sync_reconciles_csv_sqlite_and_material_snapshot(self):
        notes_path, comments_path = self._configure_csv("all-stores")
        note = {"noteId": "allstores123", "title": "全部本地同步", "content": "正文",
                "url": "https://www.xiaohongshu.com/explore/allstores123", "detailRead": True}
        comment = {"commentId": "allstores-comment", "author": "用户", "content": "没有变化",
                   "publishedAt": "08-27", "commentLevel": 1, "likeCount": 1}
        self.store.confirm(note)
        folder = Path(self.tmp.name) / "posts_materials" / "allstores123"
        folder.mkdir(parents=True)
        with self.store._session() as db:
            db.execute("UPDATE notes SET media_dir=?,source='existing_xlsx',pull_status='synced' WHERE note_id=?",
                       (str(folder), note["noteId"]))
        self.store.upsert_comments({"noteId": note["noteId"], "comments": [comment],
                                    "expectedCount": 1, "status": "likely_complete"})
        self.store._sync_pull_to_xlsx(note, [comment], {"folder": str(folder), "files": []})
        headers, _rows = self.store._read_csv_table(comments_path, COMMENT_CSV_HEADERS)
        self.store._replace_csv_table(comments_path, headers, [], "simulate-drift")

        result = self.store.sync_comment_snapshot({
            "noteId": note["noteId"], "note": note, "comments": [comment],
            "expectedCount": 1, "status": "likely_complete",
        })

        self.assertTrue(result["consistencyVerified"])
        self.assertEqual(0, result["newCount"])
        csv_rows = self._csv_rows(comments_path, COMMENT_CSV_HEADERS)
        self.assertEqual(["allstores-comment"], [row["笔记评论ID"] for row in csv_rows])
        self.assertEqual("allstores123", csv_rows[0]["笔记ID"])
        material = json.loads((folder / "comments.json").read_text(encoding="utf-8"))
        self.assertEqual(["allstores-comment"], [row["commentId"] for row in material])
        health = self.store.data_health()
        self.assertTrue(health["summary"]["relationshipsConsistent"])

    def test_wps_gb18030_csv_is_normalized_without_row_loss(self):
        notes_path = Path(self.tmp.name) / "品牌_笔记总表.csv"
        self.store.configure_data_files(notes_path)
        self.store._replace_csv_table(notes_path, NOTE_CSV_HEADERS, [], "test")
        comments_path = Path(self.tmp.name) / "品牌_评论总表.csv"
        with comments_path.open("w", encoding="gb18030", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=COMMENT_CSV_HEADERS, lineterminator="\r\n")
            writer.writeheader()
            writer.writerow({"笔记评论ID": "gb18030-comment", "用户昵称": "测试用户", "评论内容": "中文评论保持完整"})

        self.store.seed_from_xlsx(notes_path)

        self.assertEqual(b"\xef\xbb\xbf", comments_path.read_bytes()[:3])
        _headers, rows = self.store._read_csv_table(comments_path, COMMENT_CSV_HEADERS)
        self.assertEqual("中文评论保持完整", rows[0]["评论内容"])
        self.assertEqual("测试用户", rows[0]["用户昵称"])

    def test_seed_migrates_existing_workbook_to_two_utf8_csv_files(self):
        from openpyxl import Workbook

        workbook_path = self.store.export_dir.parent / "legacy-no-access-column.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "sheet1_笔记总表"
        sheet.append(["笔记url", "笔记标题", "笔记ID"])
        sheet.append([
            "https://www.xiaohongshu.com/explore/schema123456",
            "旧版表格帖子",
            "schema123456",
        ])
        workbook.create_sheet("sheet2_评论总表").append(["笔记评论ID"])
        audit = workbook.create_sheet("旧版审核记录")
        audit.append(["笔记ID", "处理结果"])
        audit.append(["schema123456", "保留"])
        workbook.save(workbook_path)
        workbook.close()

        self.store.seed_from_xlsx(workbook_path)
        notes_path, comments_path = self.store._csv_paths()
        headers, rows = self.store._read_csv_table(notes_path, NOTE_CSV_HEADERS)
        self.assertFalse(workbook_path.exists())
        self.assertTrue(comments_path.exists())
        self.assertIn("访问状态", headers)
        self.assertEqual("schema123456", rows[0]["笔记ID"])
        self.assertEqual(b"\xef\xbb\xbf", notes_path.read_bytes()[:3])
        with self.store._session() as db:
            archived = db.execute("SELECT COUNT(*) FROM legacy_sheet_archive WHERE sheet_name='旧版审核记录'").fetchone()[0]
        self.assertEqual(2, archived)

    def test_batch_access_status_writes_open_review_and_unreachable(self):
        notes_path, _comments_path = self._configure_csv("access-batch")
        notes = [
            {"noteId": "accessopen123", "title": "可打开帖子", "content": "正文"},
            {"noteId": "accessreview123", "title": "待复核帖子", "content": "正文"},
            {"noteId": "accessgone123", "title": "确认删除帖子", "content": "正文"},
        ]
        for note in notes:
            self.store.confirm(note)
            self.store._sync_pull_to_xlsx(note, [], {"folder": "", "files": []})

        result = self.store.set_note_access_statuses({"items": [
            {"noteId": "accessopen123", "status": "ok"},
            {"noteId": "accessreview123", "status": "check_failed", "error": "token 过期"},
            {"noteId": "accessgone123", "status": "unreachable", "error": "两条明确删除证据"},
        ]})
        self.assertTrue(result["ok"])
        self.assertEqual({"ok": 1, "check_failed": 1, "unreachable": 1}, result["byStatus"])

        excel_states = {row["笔记ID"]: row["访问状态"]
                        for row in self._csv_rows(notes_path, NOTE_CSV_HEADERS)}
        self.assertEqual("可打开", excel_states["accessopen123"])
        self.assertEqual("待复核", excel_states["accessreview123"])
        self.assertEqual("打不开", excel_states["accessgone123"])
        with self.store._session() as db:
            db_states = {
                row["note_id"]: row["access_status"]
                for row in db.execute(
                    "SELECT note_id,access_status FROM notes WHERE note_id LIKE 'access%'"
                ).fetchall()
            }
        self.assertEqual("ok", db_states["accessopen123"])
        self.assertEqual("check_failed", db_states["accessreview123"])
        self.assertEqual("unreachable", db_states["accessgone123"])

    def test_data_overview_filters_all_note_and_comment_fields_on_verified_snapshot(self):
        notes_path, _comments_path = self._configure_csv("data-overview")
        notes = [
            {"noteId": "overviewnote123", "url": "https://www.xiaohongshu.com/explore/overviewnote123",
             "title": "价格反馈", "content": "正文甲", "author": "作者甲", "detailRead": True},
            {"noteId": "overviewnote456", "url": "https://www.xiaohongshu.com/explore/overviewnote456",
             "title": "使用体验", "content": "正文乙", "author": "作者乙", "detailRead": True},
        ]
        comments = [
            {"commentId": "overview-comment-a", "author": "用户甲", "content": "价格太贵", "likeCount": 8},
            {"commentId": "overview-comment-b", "author": "用户乙", "content": "体验不错", "likeCount": 2},
        ]
        for note, comment in zip(notes, comments):
            self.store.confirm(note)
            self.store.upsert_comments({"noteId": note["noteId"], "comments": [comment], "status": "likely_complete"})
            self.store._sync_pull_to_xlsx(note, [comment], {"folder": "", "files": []})
        thread_note = {
            "noteId": "overviewthread123", "url": "https://www.xiaohongshu.com/explore/overviewthread123",
            "title": "评论线程", "content": "线程正文", "author": "线程作者", "detailRead": True,
        }
        thread_comments = [
            {"commentId": "overview-thread-reply-b", "parentCommentId": "overview-thread-root",
             "author": "回复乙", "content": "第二条回复", "publishedAt": "09-02", "likeCount": 2},
            {"commentId": "overview-thread-root", "author": "一级用户", "content": "一级评论正文",
             "publishedAt": "09-01", "likeCount": 5},
            {"commentId": "overview-thread-reply-a", "parentCommentId": "overview-thread-root",
             "author": "回复甲", "content": "第一条回复", "publishedAt": "09-01", "likeCount": 3},
        ]
        self.store.confirm(thread_note)
        self.store.upsert_comments({
            "noteId": thread_note["noteId"], "comments": thread_comments,
            "expectedCount": 3, "status": "likely_complete",
        })
        self.store._sync_pull_to_xlsx(thread_note, thread_comments, {"folder": "", "files": []})
        schema = self.store.data_overview_schema()
        self.assertTrue(schema["queryReady"])
        self.assertTrue(schema["health"]["summary"]["relationshipsConsistent"])
        note_fields = {item["key"] for item in schema["datasets"]["notes"]["fields"]}
        comment_fields = {item["key"] for item in schema["datasets"]["comments"]["fields"]}
        self.assertTrue({"note_id", "payload_json", "active_comment_count", "source_like_count"}.issubset(note_fields))
        self.assertTrue(next(item for item in schema["datasets"]["notes"]["fields"]
                             if item["key"] == "ignore_status")["defaultVisible"])
        self.assertTrue({"comment_id", "content", "post__title", "post__payload_json"}.issubset(comment_fields))
        comment_field_rows = schema["datasets"]["comments"]["fields"]
        self.assertEqual(
            ["comment_id", "note_id", "thread_root_content", "content", "author", "published_at", "like_count",
             "comment_level", "analysis_is_negative", "negative_type", "comment_status", "comment_type",
             "post__url", "post__title", "post__post_status"],
            [item["key"] for item in comment_field_rows if item["defaultVisible"]],
        )
        self.assertTrue(next(item for item in comment_field_rows if item["key"] == "comment_status")["suggestValues"])

        note_result = self.store.query_data_overview({
            "dataset": "notes", "snapshotToken": schema["snapshotToken"],
            "fields": ["note_id", "title", "active_comment_count"],
            "filter": {"logic": "and", "children": [
                {"field": "title", "operator": "contains", "value": "价格"},
                {"field": "active_comment_count", "operator": "gte", "value": 1},
            ]},
        })
        self.assertEqual(1, note_result["total"])
        self.assertEqual("overviewnote123", note_result["rows"][0]["note_id"])
        comment_result = self.store.query_data_overview({
            "dataset": "comments", "snapshotToken": schema["snapshotToken"],
            "fields": ["comment_id", "content", "post__title"],
            "filter": {"logic": "and", "children": [
                {"field": "like_count", "operator": "gt", "value": 5},
                {"field": "post__author", "operator": "eq", "value": "作者甲"},
            ]},
        })
        self.assertEqual(1, comment_result["total"])
        self.assertEqual("价格反馈", comment_result["rows"][0]["post__title"])
        thread_result = self.store.query_data_overview({
            "dataset": "comments", "snapshotToken": schema["snapshotToken"], "groupThreads": True,
            "fields": ["comment_id", "thread_root_id", "thread_root_content", "thread_root_author", "content"],
            "filter": {"logic": "and", "children": [
                {"field": "note_id", "operator": "eq", "value": thread_note["noteId"]},
            ]},
            "pageSize": 10,
        })
        self.assertEqual(
            ["overview-thread-root", "overview-thread-reply-a", "overview-thread-reply-b"],
            [row["comment_id"] for row in thread_result["rows"]],
        )
        self.assertEqual({"一级评论正文"}, {row["thread_root_content"] for row in thread_result["rows"]})
        value_result = self.store.data_overview_values({
            "dataset": "comments", "snapshotToken": schema["snapshotToken"],
            "field": "comment_type", "limit": 20,
        })
        self.assertTrue(value_result["consistentSnapshot"])
        self.assertTrue({"一级评论", "二级回复"}.issubset({item["value"] for item in value_result["values"]}))
        with self.store._session() as db:
            db.execute("UPDATE notes SET last_seen_at=? WHERE note_id=?", ("2099-01-01", notes[0]["noteId"]))
        with self.assertRaisesRegex(ValueError, "数据已变化"):
            self.store.query_data_overview({"dataset": "notes", "snapshotToken": schema["snapshotToken"]})

    def test_data_overview_note_scope_excludes_discovery_and_includes_ignored(self):
        self._configure_csv("data-overview-scope")
        business = {
            "noteId": "overview-business-1", "url": "https://www.xiaohongshu.com/explore/overview-business-1",
            "title": "业务总表帖子", "content": "正文", "author": "作者", "detailRead": True,
        }
        confirmed = {
            "noteId": "overview-confirmed-1", "url": "https://www.xiaohongshu.com/explore/overview-confirmed-1",
            "title": "待同步帖子", "content": "正文", "author": "作者", "detailRead": True,
        }
        ignored = {
            "noteId": "overview-ignored-1", "url": "https://www.xiaohongshu.com/explore/overview-ignored-1",
            "title": "已忽略帖子", "content": "正文", "author": "作者", "detailRead": True,
        }
        discovery = {
            "noteId": "overview-discovery-1", "url": "https://www.xiaohongshu.com/explore/overview-discovery-1",
            "title": "仅发现未入库", "content": "正文", "author": "作者",
        }
        self.store.confirm(business)
        self.store._sync_pull_to_xlsx(business, [], {"folder": "", "files": []})
        self.store.confirm(confirmed)
        self.store.confirm(ignored)
        self.store.ignore({"noteId": ignored["noteId"]})
        self.store.confirm(discovery)
        with self.store._session() as db:
            db.execute("UPDATE notes SET source='existing_xlsx',pull_status='synced',status='known' WHERE note_id=?",
                       (business["noteId"],))
            db.execute("UPDATE notes SET source='dom',pull_status='not_started',status='new' WHERE note_id=?",
                       (discovery["noteId"],))

        schema = self.store.data_overview_schema()
        self.assertTrue(schema["queryReady"])
        self.assertEqual(3, schema["datasets"]["notes"]["total"])
        self.assertEqual(4, schema["datasets"]["notes"]["recordTotal"])
        self.assertEqual(1, schema["datasets"]["notes"]["ignoredTotal"])
        self.assertEqual(1, schema["datasets"]["notes"]["excludedDiscoveryTotal"])
        result = self.store.query_data_overview({
            "dataset": "notes", "snapshotToken": schema["snapshotToken"],
            "fields": ["note_id", "ignore_status"], "pageSize": 20,
        })
        self.assertEqual(
            {business["noteId"], confirmed["noteId"], ignored["noteId"]},
            {row["note_id"] for row in result["rows"]},
        )
        self.assertEqual("已忽略", next(row["ignore_status"] for row in result["rows"]
                                          if row["note_id"] == ignored["noteId"]))

    def test_purge_untracked_discoveries_preserves_formal_and_ignored_records(self):
        self._configure_csv("data-overview-purge")
        business = {
            "noteId": "purge-business-1", "url": "https://www.xiaohongshu.com/explore/purge-business-1",
            "title": "正式记录", "content": "正文", "author": "作者", "detailRead": True,
        }
        ignored = {
            "noteId": "purge-ignored-1", "url": "https://www.xiaohongshu.com/explore/purge-ignored-1",
            "title": "忽略记录", "content": "正文", "author": "作者", "detailRead": True,
        }
        discovery = {
            "noteId": "purge-discovery-1", "url": "https://www.xiaohongshu.com/explore/purge-discovery-1",
            "title": "仅发现记录", "content": "正文", "author": "作者",
        }
        self.store.confirm(business)
        self.store._sync_pull_to_xlsx(business, [], {"folder": "", "files": []})
        self.store.confirm(ignored)
        self.store.ignore({"noteId": ignored["noteId"]})
        self.store.confirm(discovery)
        with self.store._session() as db:
            db.execute("UPDATE notes SET source='existing_xlsx',pull_status='synced',status='known' WHERE note_id=?",
                       (business["noteId"],))
            db.execute("UPDATE notes SET source='dom',pull_status='not_started',status='new' WHERE note_id=?",
                       (discovery["noteId"],))
            db.execute("""INSERT INTO ai_jobs
                       (target_type,target_id,priority,status,attempts,available_at,created_at,updated_at)
                       VALUES ('relevance',?,50,'queued',0,0,'2026-09-02','2026-09-02')""",
                       (discovery["noteId"],))
            db.execute("""INSERT INTO ai_analysis_records(target_type,target_id,status,created_at)
                       VALUES ('relevance',?,'completed','2026-09-02')""", (discovery["noteId"],))

        schema = self.store.data_overview_schema()
        self.assertTrue(schema["queryReady"])
        preview = self.store.purge_untracked_discoveries({
            "snapshotToken": schema["snapshotToken"], "dryRun": True,
        })
        self.assertEqual(1, preview["candidateCount"])
        self.assertEqual({"new": 1}, preview["byStatus"])
        result = self.store.purge_untracked_discoveries({
            "snapshotToken": schema["snapshotToken"], "hardDeleteConfirmed": True,
            "confirmation": preview["confirmation"],
        })
        self.assertEqual(1, result["deletedCount"])
        self.assertEqual(2, result["deletedLinkedDatabaseRecords"])
        with self.store._session() as db:
            self.assertEqual(
                {business["noteId"], ignored["noteId"]},
                {row[0] for row in db.execute("SELECT note_id FROM notes").fetchall()},
            )
            self.assertEqual(0, db.execute(
                "SELECT COUNT(*) FROM ai_jobs WHERE target_id=?", (discovery["noteId"],)
            ).fetchone()[0])
            self.assertEqual(0, db.execute(
                "SELECT COUNT(*) FROM ai_analysis_records WHERE target_id=?", (discovery["noteId"],)
            ).fetchone()[0])
        refreshed = self.store.data_overview_schema()
        self.assertEqual(2, refreshed["datasets"]["notes"]["recordTotal"])
        self.assertEqual(2, refreshed["datasets"]["notes"]["total"])
        self.assertEqual(0, refreshed["datasets"]["notes"]["excludedDiscoveryTotal"])

    def test_data_overview_permanent_delete_cascades_every_local_store(self):
        notes_path, comments_path = self._configure_csv("data-overview-delete")
        note = {
            "noteId": "overview-delete-1", "url": "https://www.xiaohongshu.com/explore/overview-delete-1",
            "title": "永久删除测试", "content": "正文", "author": "作者", "detailRead": True,
        }
        comments = [
            {"commentId": "overview-delete-root", "author": "一级用户", "content": "一级评论", "commentLevel": 1},
            {"commentId": "overview-delete-reply", "parentCommentId": "overview-delete-root",
             "author": "回复用户", "content": "二级回复", "commentLevel": 2},
            {"commentId": "overview-delete-keep", "author": "保留用户", "content": "保留评论", "commentLevel": 1},
        ]
        self.store.confirm(note)
        self.store.upsert_comments({"noteId": note["noteId"], "comments": comments,
                                    "expectedCount": 3, "status": "likely_complete"})
        folder = self.store._media_root() / f"永久删除测试__{note['noteId']}"
        folder.mkdir(parents=True)
        (folder / "image-01.jpg").write_bytes(b"image")
        (folder / "note.json").write_text(json.dumps({
            **note, "postStatus": "存在", "isDeleted": False, "accessStatus": "ok"
        }, ensure_ascii=False), encoding="utf-8")
        (folder / "帖子正文.txt").write_text("永久删除测试\n\n正文", encoding="utf-8")
        media_comments = self.store._comments_as_api(note["noteId"])
        (folder / "comments.json").write_text(json.dumps(media_comments, ensure_ascii=False), encoding="utf-8")
        media = {"folder": str(folder), "files": ["image-01.jpg", "note.json", "comments.json", "帖子正文.txt"]}
        self.store._sync_pull_to_xlsx(note, comments, media)
        with self.store._session() as db:
            db.execute("""UPDATE notes SET source='existing_xlsx',pull_status='synced',status='known',
                       media_dir=?,media_file_count=4,access_status='ok' WHERE note_id=?""",
                       (str(folder), note["noteId"]))

        token = self.store._data_overview_snapshot_token()
        self.store._data_overview_approved_tokens[token] = time.time()
        deleted_comments = self.store.delete_data_overview_records({
            "dataset": "comments", "ids": ["overview-delete-root"], "snapshotToken": token,
            "hardDeleteConfirmed": True, "confirmation": "DELETE:comments:1",
        })
        self.assertEqual(2, deleted_comments["deletedCount"])
        self.assertEqual(1, deleted_comments["cascadeDeletedCount"])
        self.assertTrue(folder.is_dir())
        self.assertEqual(["overview-delete-keep"], [
            item["commentId"] for item in json.loads((folder / "comments.json").read_text(encoding="utf-8"))
        ])
        self.assertEqual(["overview-delete-keep"], [
            row["笔记评论ID"] for row in self._csv_rows(comments_path, COMMENT_CSV_HEADERS)
        ])
        with self.store._session() as db:
            self.assertEqual(["overview-delete-keep"], [row[0] for row in db.execute(
                "SELECT comment_id FROM comments WHERE note_id=? ORDER BY comment_id", (note["noteId"],)
            ).fetchall()])
            self.assertEqual(1, db.execute(
                "SELECT comment_count_collected FROM notes WHERE note_id=?", (note["noteId"],)
            ).fetchone()[0])

        token = self.store._data_overview_snapshot_token()
        self.store._data_overview_approved_tokens[token] = time.time()
        deleted_note = self.store.delete_data_overview_records({
            "dataset": "notes", "ids": [note["noteId"]], "snapshotToken": token,
            "hardDeleteConfirmed": True, "confirmation": "DELETE:notes:1",
        })
        self.assertEqual(1, deleted_note["deletedCount"])
        self.assertEqual(1, deleted_note["deletedCommentCount"])
        self.assertFalse(folder.exists())
        self.assertFalse(any(row["笔记ID"] == note["noteId"]
                             for row in self._csv_rows(notes_path, NOTE_CSV_HEADERS)))
        self.assertFalse(any(row["笔记ID"] == note["noteId"]
                             for row in self._csv_rows(comments_path, COMMENT_CSV_HEADERS)))
        with self.store._session() as db:
            self.assertEqual(0, db.execute("SELECT COUNT(*) FROM notes WHERE note_id=?", (note["noteId"],)).fetchone()[0])
            self.assertEqual(0, db.execute("SELECT COUNT(*) FROM comments WHERE note_id=?", (note["noteId"],)).fetchone()[0])

    def test_data_overview_snapshot_tracks_material_json_changes(self):
        self._configure_csv("data-overview-material-token")
        note = {
            "noteId": "overviewmaterial123", "url": "https://www.xiaohongshu.com/explore/overviewmaterial123",
            "title": "素材快照", "content": "正文", "author": "作者", "detailRead": True,
        }
        self.store.confirm(note)
        folder = Path(self.tmp.name) / "materials" / note["noteId"]
        folder.mkdir(parents=True)
        (folder / "note.json").write_text('{"revision":1}', encoding="utf-8")
        (folder / "comments.json").write_text('[]', encoding="utf-8")
        (folder / "帖子正文.txt").write_text("正文", encoding="utf-8")
        with self.store._session() as db:
            db.execute("UPDATE notes SET media_dir=? WHERE note_id=?", (str(folder), note["noteId"]))
        first = self.store._data_overview_snapshot_token()
        (folder / "note.json").write_text('{"revision":222}', encoding="utf-8")
        second = self.store._data_overview_snapshot_token()
        self.assertNotEqual(first, second)

    def test_data_overview_extension_page_and_workbench_launcher_exist(self):
        extension = Path(__file__).resolve().parent.parent / "extension"
        page = (extension / "data-overview.html").read_text(encoding="utf-8")
        script = (extension / "data-overview.js").read_text(encoding="utf-8")
        worker = (extension / "service-worker.js").read_text(encoding="utf-8")
        panel = (extension / "sidepanel.html").read_text(encoding="utf-8")
        self.assertIn('id="openDataOverview"', panel)
        self.assertIn('id="filterRows"', page)
        self.assertIn('id="fieldOptions"', page)
        self.assertIn('id="infiniteSentinel"', page)
        self.assertIn('id="pageFindInput"', page)
        self.assertIn('id="columnFilterPopover"', page)
        self.assertIn('id="deleteSelected"', page)
        self.assertIn('id="deleteDialog"', page)
        self.assertNotIn('id="previousPage"', page)
        self.assertIn('type: "getDataOverviewSchema"', script)
        self.assertIn("IntersectionObserver", script)
        self.assertIn("loadNextBatch", script)
        self.assertIn("openColumnFilter", script)
        self.assertIn("applyColumnFilter", script)
        self.assertIn("performPermanentDelete", script)
        self.assertIn('type: "deleteDataOverviewRecords"', script)
        self.assertGreaterEqual(script.count("await loadSchema({ preserveQuery: true })"), 2)
        self.assertIn('message.type === "queryDataOverview"', worker)
        self.assertIn('message.type === "getDataOverviewValues"', worker)
        self.assertIn('message.type === "deleteDataOverviewRecords"', worker)
        self.assertIn("relationshipsConsistent", script)
        self.assertIn("queryPending", script)

    def test_summary_combines_note_and_comments(self):
        settings = self.store.ai_settings._raw()
        settings["api_key_dpapi"] = "test"
        self.store.ai_settings.get = lambda include_secret=False: {
            "configured": True, "model": "deepseek-v4-flash", "max_tokens": 1800,
            "api_key": "test", "base_url": "https://api.deepseek.com", "timeout_seconds": 45,
            "temperature": 0.1, "thinking_mode": "disabled"
        }
        result = self.store.summarize_note({"note": {"noteId": "abcdef123456", "title": "标题", "content": "正文内容"},
                                            "comments": [{"content": "代表评论", "author": "用户"}]})
        self.assertTrue(result["ok"])
        self.assertEqual(1, result["commentCount"])
        cached = self.store.note_summary("abcdef123456")
        self.assertTrue(cached["found"])

    def test_reply_suggestion_uses_manual_send_workflow(self):
        self.store.ai_client = FakeReplyAI()
        self.store.ai_settings.get = lambda include_secret=False: {
            "configured": True, "model": "deepseek-v4-flash", "max_tokens": 1800,
            "api_key": "test", "base_url": "https://api.deepseek.com", "timeout_seconds": 45,
            "temperature": 0.1, "thinking_mode": "disabled"
        }
        result = self.store.suggest_comment_reply({
            "note": {"noteId": "abcdef123456", "title": "标题", "content": "正文"},
            "comments": [{"commentId": "comment-1", "content": "价格是多少", "author": "用户"}],
            "targetComment": {"commentId": "comment-1", "content": "价格是多少", "author": "用户"},
            "persona": "brand"
        })
        self.assertTrue(result["ok"])
        self.assertEqual(3, len(result["suggestion"]["candidates"]))
        self.assertIn("价格", result["suggestion"]["reply"])
        self.assertEqual("价格疑问", result["suggestion"]["intent"])
        self.assertTrue(all(item["riskLevel"] != "high" for item in result["suggestion"]["candidates"]))
        with self.store._session() as db:
            self.assertEqual(3, db.execute("SELECT COUNT(*) FROM reply_generation_history").fetchone()[0])

    def test_reply_policy_flags_template_marketing_and_mismatch(self):
        intent, sentiment = classify_reply_context("我今天被拉着推销，感觉很不舒服")
        self.assertEqual(("推销投诉", "负面"), (intent, sentiment))
        audit = audit_reply_candidate(
            "感谢宝宝认可与喜爱，未来2-3个月会有更多优惠活动，欢迎来店体验。",
            "我今天被拉着推销，感觉很不舒服", "brand", []
        )
        self.assertEqual("high", audit["riskLevel"])
        self.assertTrue(any("语义不匹配" in item for item in audit["riskNotes"]))

    def test_reply_policy_flags_high_similarity(self):
        reply = "理解你对价格的顾虑。你发下产品全名，我们按门店公示价帮你核对。"
        audit = audit_reply_candidate(reply, "这个多少钱", "brand", [reply])
        self.assertEqual("high", audit["riskLevel"])
        self.assertGreaterEqual(audit["similarityScore"], 0.99)

    def test_operations_center_health_changes_watchlist_and_weekly_report(self):
        from datetime import datetime
        from openpyxl import load_workbook

        workbook_path = self.store.export_dir.parent / "operations-master.xlsx"
        self.store.seed_xlsx_path = workbook_path
        self.store._ensure_seed_workbook(workbook_path)
        note = {
            "noteId": "operations123456", "title": "重点观察测试帖", "author": "测试作者",
            "content": "有人反馈价格和推销问题", "url": "https://www.xiaohongshu.com/explore/operations123456",
            "keyword": "samplebrand", "detailRead": True, "likeCount": 1,
        }
        self.store.confirm(note)
        original = {"commentId": "operations-comment-1", "author": "用户甲", "content": "价格多少", "publishedAt": "08-26"}
        self.store.upsert_comments({
            "noteId": note["noteId"], "comments": [original], "expectedCount": 1, "status": "likely_complete"
        })
        self.store._sync_pull_to_xlsx(note, [original], {"folder": "", "files": []})

        watched = self.store.set_watchlist({
            "noteId": note["noteId"], "watched": True, "priority": "high", "reason": "价格投诉"
        })
        self.assertTrue(watched["watched"])
        self.assertEqual(1, self.store.list_watchlist()["count"])
        self.assertTrue(self.store.note_status(note["noteId"])["watched"])

        run = self.store.start_sync_run({"runType": "batch", "totalNotes": 1})
        added = {"commentId": "operations-comment-2", "author": "用户乙", "content": "被拉着推销", "publishedAt": "08-26"}
        updated_note = {**note, "likeCount": 2}
        synced = self.store.sync_comment_snapshot({
            "noteId": note["noteId"], "note": updated_note, "comments": [original, added],
            "expectedCount": 2, "status": "likely_complete", "runId": run["runId"],
        })
        self.assertEqual(1, synced["newCount"])
        self.assertEqual(2, synced["changeEventCount"])
        self.store.finish_sync_run({
            "runId": run["runId"], "status": "completed", "processedNotes": 1,
            "changedNotes": 1, "newComments": 1,
        })
        changes = self.store.list_change_events()
        self.assertEqual(2, changes["unread"])
        self.assertEqual({"comment_added", "note_fields_changed"}, {item["event_type"] for item in changes["events"]})
        self.assertEqual(2, self.store.acknowledge_change_events({"all": True})["updated"])
        self.assertEqual(0, self.store.list_change_events()["unread"])

        with self.store._session() as db:
            db.execute("UPDATE notes SET comment_count_collected=99 WHERE note_id=?", (note["noteId"],))
        health = self.store.data_health()
        mismatch = next(item for item in health["issues"] if item["id"] == "comment_count_mismatch")
        self.assertTrue(mismatch["repairable"])
        repaired = self.store.repair_data_health({})
        self.assertTrue(repaired["ok"])
        self.assertFalse(any(item["id"] == "comment_count_mismatch" for item in repaired["health"]["issues"]))

        today = datetime.now().astimezone().date().isoformat()
        report = self.store.generate_weekly_report({"startDate": today, "endDate": today})
        self.assertTrue(Path(report["xlsxPath"]).is_file())
        self.assertTrue(Path(report["htmlPath"]).is_file())
        self.assertEqual(1, report["summary"]["newComments"])
        workbook = load_workbook(report["xlsxPath"], read_only=True, data_only=True)
        self.assertEqual({"周报总览", "新增帖子", "同步变化", "重点观察"}, set(workbook.sheetnames))
        workbook.close()
        self.assertTrue(self.store.latest_weekly_report()["found"])

    def test_restart_cancels_legacy_sentiment_queue_and_disables_auto_flags(self):
        note = {"noteId": "sentiment123456", "title": "旧情绪任务", "content": "正文"}
        self.store.confirm(note)
        self.store.upsert_comments({
            "noteId": note["noteId"],
            "comments": [{"commentId": "sentiment-comment-1", "author": "用户", "content": "评论"}],
            "expectedCount": 1,
            "status": "likely_complete",
        })
        self.store.ai_settings.save({"auto_analyze_posts": True, "auto_analyze_comments": True})
        with self.store._session() as db:
            db.execute("UPDATE notes SET ai_analysis_status='queued' WHERE note_id=?", (note["noteId"],))
            db.execute("UPDATE comments SET ai_analysis_status='analyzing' WHERE comment_id='sentiment-comment-1'")
            db.execute("""INSERT INTO ai_jobs
                       (target_type,target_id,priority,status,attempts,available_at,created_at,updated_at)
                       VALUES ('note',?,10,'queued',0,0,'2026-08-25','2026-08-25')""", (note["noteId"],))
            db.execute("""INSERT INTO ai_jobs
                       (target_type,target_id,priority,status,attempts,available_at,created_at,updated_at)
                       VALUES ('comment','sentiment-comment-1',10,'analyzing',0,0,'2026-08-25','2026-08-25')""")

        reopened = MonitorStore(self.store.db_path, self.store.export_dir, ai_client=FakeAI())
        settings = reopened.ai_settings.get(False)
        self.assertFalse(settings["auto_analyze_posts"])
        self.assertFalse(settings["auto_analyze_comments"])
        with reopened._session() as db:
            self.assertEqual(0, db.execute(
                "SELECT COUNT(*) FROM ai_jobs WHERE status IN ('queued','analyzing')"
            ).fetchone()[0])
            self.assertEqual("not_analyzed", db.execute(
                "SELECT ai_analysis_status FROM notes WHERE note_id=?", (note["noteId"],)
            ).fetchone()[0])
            self.assertEqual("not_analyzed", db.execute(
                "SELECT ai_analysis_status FROM comments WHERE comment_id='sentiment-comment-1'"
            ).fetchone()[0])


if __name__ == "__main__":
    unittest.main()
